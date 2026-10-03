"""Generate the feature-catalog views from the single YAML source, and check it.

The whole product's feature list lives in one file,
``docs/01-requirements/features/features.yaml``. This script renders it:

- ``generate``: writes ``README.md`` (progress per domain and surface, the
  domain index, legend, non-functional requirements, old-code mapping),
  ``domains/<domain>.md`` (one checklist per domain, then every feature in
  detail), and ``features.xlsx`` (in Vietnamese, for business readers: a
  "Tổng quan" sheet indexing every domain with links, one sheet per domain,
  and one sheet of non-functional requirements).
- ``check``: fails when the source breaks a rule (unknown role, surface,
  status, priority, release or offer; a duplicate code; a dependency on a
  missing feature or a cycle; a table that is not in the DBML; a decision ID
  that is not in the decision log; a future.md item that does not exist; a
  missing Vietnamese text) or when any generated view is stale.

Run from the repository root::

    uv run --project backend --with pyyaml --with openpyxl python \\
        .claude/skills/feature-catalog/scripts/feature_catalog.py generate
    uv run --project backend --with pyyaml --with openpyxl python \\
        .claude/skills/feature-catalog/scripts/feature_catalog.py check

Limitations: ``Data IN-n`` / ``Data OUT-n`` source references are not checked
against the phase-1 data sheet, and nothing checks a status against the code;
the backend status is synced by hand.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import yaml

# The script lives at <root>/.claude/skills/feature-catalog/scripts/, so the
# repository root is four levels up.
REPO_ROOT = Path(__file__).resolve().parents[4]
CATALOG_DIR = REPO_ROOT / "docs" / "01-requirements" / "features"
SOURCE_PATH = CATALOG_DIR / "features.yaml"
README_PATH = CATALOG_DIR / "README.md"
DOMAINS_DIR = CATALOG_DIR / "domains"
XLSX_PATH = CATALOG_DIR / "features.xlsx"
# Other documents the catalog points into; check verifies the references.
DBML_PATH = (
    REPO_ROOT / "docs" / "01-requirements" / "domain-model" / "domain-model.dbml"
)
DECISION_LOG_PATH = REPO_ROOT / "docs" / "05-decisions" / "decision-log.md"
FUTURE_PATHS = (
    REPO_ROOT / "docs" / "01-requirements" / "future.md",
    REPO_ROOT / "docs" / "01-requirements" / "future-resolved.md",
)

GENERATED_MARKER = "<!-- GENERATED from features.yaml"
GENERATED_HEADER = (
    f"{GENERATED_MARKER} by the feature-catalog skill. "
    "Edit the source, then regenerate; never edit this file by hand. -->\n\n"
)

# Where a feature is built, in display order. A feature lists only the
# surfaces it needs.
SURFACES = ("backend", "app", "portal")
SURFACE_LABELS = {"backend": "Backend", "app": "App", "portal": "Portal"}
VI_SURFACE_LABELS = {"backend": "Backend", "app": "App", "portal": "Portal web"}
STATUSES = ("todo", "doing", "done")
STATUS_MARKS = {"todo": "⬜", "doing": "🚧", "done": "✅"}
VI_STATUS_LABELS = {"todo": "Chưa làm", "doing": "Đang làm", "done": "Xong"}
VI_NOT_NEEDED = "Không cần"
PRIORITIES = ("Must", "Should", "Could")
VI_PRIORITY_LABELS = {"Must": "Bắt buộc", "Should": "Nên có", "Could": "Có thì tốt"}
RELEASES = ("P1.0", "P1.1", "P1.5", "P2")
OFFER_LABELS = {
    "internal": ("Internal only", "Chỉ nội bộ (không bán)"),
    "included": ("Included in every plan", "Có trong mọi gói"),
    "tbd": ("To be priced", "Chưa định giá"),
}

FEATURE_REQUIRED_TEXT = (
    "code",
    "name",
    "name_vi",
    "what",
    "what_vi",
    "value",
    "value_vi",
    "priority",
    "release",
    "offer",
)
FEATURE_REQUIRED_LISTS = ("users", "capabilities", "sources")
FEATURE_OPTIONAL_LISTS = (
    "depends_on",
    "touches",
    "related_tables",
    "old_codes",
    "questions",
)
FEATURE_KEYS = {
    *FEATURE_REQUIRED_TEXT,
    *FEATURE_REQUIRED_LISTS,
    *FEATURE_OPTIONAL_LISTS,
    "surfaces",
}
DOMAIN_KEYS = {"prefix", "key", "name", "name_vi", "summary", "summary_vi"}
DOMAIN_OPTIONAL_KEYS = {"backend_domain", "sheet_vi"}

CODE_PATTERN = re.compile(r"^([A-Z]{3})-(\d{2})$")
OLD_CODE_PATTERN = re.compile(r"^F-[A-K]\d$")
NFR_CODE_PATTERN = re.compile(r"^NF-\d{2}$")
# Every source reference starts with one of these; the Vietnamese prefix is
# what the workbook shows.
SOURCE_PREFIXES = {
    "PRD ": "PRD ",
    "Data ": "Bảng dữ liệu ",
    "Phase 2 ": "Phase 2 ",
    "Prerequisite ": "Điều kiện tiên quyết ",
    "Decision ": "Quyết định ",
    "future.md ": "future.md ",
    "Design review ": "Rà soát thiết kế ",
    "NF-": "NF-",
}


class CatalogError(Exception):
    """The source file cannot be parsed into a catalog at all."""


@dataclass(frozen=True)
class Feature:
    """One feature of the catalog, as written in the source.

    Attributes:
        code: Stable code ``<PREFIX>-<NN>``.
        name, name_vi: Short name in English and Vietnamese.
        what, what_vi: What the feature does for its users.
        value, value_vi: Why it matters (business effect).
        users: Role codes that use or receive the feature.
        capabilities: (English, Vietnamese) pairs breaking the feature down.
        surfaces: Surface -> status, only for the surfaces the feature needs.
        priority, release, offer: Planning and pricing labels.
        sources: Where the feature comes from.
        depends_on: Codes of features that must exist first.
        touches: Other backend domains the feature reads from or writes to.
        related_tables: DBML tables it uses (empty until the database review).
        old_codes: Codes in the old feature-list.md.
        questions: Open (English, Vietnamese) question pairs.
    """

    code: str
    name: str
    name_vi: str
    what: str
    what_vi: str
    value: str
    value_vi: str
    users: tuple[str, ...]
    capabilities: tuple[tuple[str, str], ...]
    surfaces: dict[str, str]
    priority: str
    release: str
    offer: str
    sources: tuple[str, ...]
    depends_on: tuple[str, ...]
    touches: tuple[str, ...]
    related_tables: tuple[str, ...]
    old_codes: tuple[str, ...]
    questions: tuple[tuple[str, str], ...]

    @property
    def prefix(self) -> str:
        """Domain prefix of the code, e.g. ``ACC``."""
        return self.code.split("-")[0]

    @property
    def is_done(self) -> bool:
        """Whether every surface the feature needs is done."""
        return all(status == "done" for status in self.surfaces.values())


@dataclass
class Domain:
    """A group of features sharing one code prefix and one page/sheet.

    Attributes:
        prefix: Three-letter code prefix, e.g. ``ACC``.
        key: File name of the domain page, e.g. ``identity``.
        name, name_vi: Title in English and Vietnamese.
        summary, summary_vi: One-paragraph scope.
        backend_domain: Backend domain that owns it, if any.
        sheet_vi: Short Vietnamese label for the workbook sheet name, when
            the prefix and ``name_vi`` exceed Excel's 31 characters.
        features: Its features, ordered by code.
    """

    prefix: str
    key: str
    name: str
    name_vi: str
    summary: str
    summary_vi: str
    backend_domain: str | None
    sheet_vi: str | None = None
    features: list[Feature] = field(default_factory=list)

    @property
    def title(self) -> str:
        """English title with the prefix, e.g. ``ACC — Accounts & access``."""
        return f"{self.prefix} — {self.name}"


@dataclass(frozen=True)
class Nfr:
    """One non-functional requirement.

    Attributes:
        code: ``NF-NN``, kept from the PRD.
        category, requirement, target: (English, Vietnamese) pairs.
        features: Codes of the features it applies to most directly.
    """

    code: str
    category: tuple[str, str]
    requirement: tuple[str, str]
    target: tuple[str, str]
    features: tuple[str, ...]


@dataclass
class Catalog:
    """The parsed catalog and every rule it breaks.

    Attributes:
        roles: Role code -> (English, Vietnamese) label.
        domains: Domains in source order.
        features: Feature code -> feature.
        nfrs: Non-functional requirements in source order.
        problems: Rule violations; ``check`` fails and ``generate`` refuses
            to write while any exist.
    """

    roles: dict[str, tuple[str, str]]
    domains: list[Domain]
    features: dict[str, Feature]
    nfrs: list[Nfr]
    problems: list[str]

    def domain_of(self, code: str) -> Domain:
        """Return the domain owning a feature code."""
        prefix = code.split("-")[0]
        return next(domain for domain in self.domains if domain.prefix == prefix)

    def dependents_of(self, code: str) -> list[str]:
        """Codes of the features that depend on ``code``, sorted."""
        return sorted(
            feature.code
            for feature in self.features.values()
            if code in feature.depends_on
        )


# ---------------------------------------------------------------------------
# Loading and validation
# ---------------------------------------------------------------------------


def _text_pair(raw: object, where: str, problems: list[str]) -> tuple[str, str]:
    """Read an ``[English, Vietnamese]`` pair, recording a problem if malformed.

    Args:
        raw: The YAML value.
        where: Location used in the problem message.
        problems: List the problem is appended to.

    Returns:
        The pair, or two empty strings when malformed.
    """
    if (
        isinstance(raw, list)
        and len(raw) == 2
        and all(isinstance(part, str) and part.strip() for part in raw)
    ):
        return (raw[0].strip(), raw[1].strip())
    problems.append(f"{where}: expected a non-empty [English, Vietnamese] pair")
    return ("", "")


def _string_list(raw: object, where: str, problems: list[str]) -> tuple[str, ...]:
    """Read a list of non-empty strings, recording a problem if malformed."""
    if raw is None:
        return ()
    if not isinstance(raw, list) or not all(
        isinstance(item, str) and item.strip() for item in raw
    ):
        problems.append(f"{where}: expected a list of non-empty strings")
        return ()
    return tuple(item.strip() for item in raw)


def _parse_feature(raw: dict[str, object], problems: list[str]) -> Feature | None:
    """Turn one raw feature mapping into a Feature, recording shape problems.

    Returns:
        The feature, or None when it lacks a usable code.
    """
    code = raw.get("code")
    if not isinstance(code, str) or not CODE_PATTERN.match(code):
        problems.append(f"feature {code!r}: code must look like ACC-01")
        return None
    for key in sorted(set(raw) - FEATURE_KEYS):
        problems.append(f"{code}: unknown field {key!r}")
    texts: dict[str, str] = {}
    for key in FEATURE_REQUIRED_TEXT:
        value = raw.get(key)
        if not isinstance(value, str) or not value.strip():
            problems.append(f"{code}: missing {key}")
            value = ""
        texts[key] = " ".join(str(value).split())
    for key in FEATURE_REQUIRED_LISTS:
        if not raw.get(key):
            problems.append(f"{code}: {key} must not be empty")
    capabilities = tuple(
        _text_pair(item, f"{code} capability {index}", problems)
        for index, item in enumerate(raw.get("capabilities") or [], start=1)
    )
    questions = tuple(
        _text_pair(item, f"{code} question {index}", problems)
        for index, item in enumerate(raw.get("questions") or [], start=1)
    )
    surfaces_raw = raw.get("surfaces")
    surfaces: dict[str, str] = {}
    if not isinstance(surfaces_raw, dict) or not surfaces_raw:
        problems.append(f"{code}: surfaces must map at least one surface to a status")
    else:
        for surface in SURFACES:
            if surface in surfaces_raw:
                surfaces[surface] = str(surfaces_raw[surface])
        for surface in sorted(set(surfaces_raw) - set(SURFACES)):
            problems.append(f"{code}: unknown surface {surface!r}")
    return Feature(
        code=code,
        name=texts["name"],
        name_vi=texts["name_vi"],
        what=texts["what"],
        what_vi=texts["what_vi"],
        value=texts["value"],
        value_vi=texts["value_vi"],
        users=_string_list(raw.get("users"), f"{code} users", problems),
        capabilities=capabilities,
        surfaces=surfaces,
        priority=texts["priority"],
        release=texts["release"],
        offer=texts["offer"],
        sources=_string_list(raw.get("sources"), f"{code} sources", problems),
        depends_on=_string_list(raw.get("depends_on"), f"{code} depends_on", problems),
        touches=_string_list(raw.get("touches"), f"{code} touches", problems),
        related_tables=_string_list(
            raw.get("related_tables"), f"{code} related_tables", problems
        ),
        old_codes=_string_list(raw.get("old_codes"), f"{code} old_codes", problems),
        questions=questions,
    )


def _read_reference_sets() -> tuple[set[str], set[str], set[str], set[str]]:
    """Collect the names other documents define, for checking references.

    Returns:
        DBML table names, DBML table-group (backend domain) names, decision
        IDs of the decision log, and future.md item numbers (open and
        resolved). A missing document yields an empty set.
    """
    dbml = DBML_PATH.read_text(encoding="utf-8") if DBML_PATH.exists() else ""
    tables = set(re.findall(r"^Table (\w+)", dbml, flags=re.MULTILINE))
    groups = set(re.findall(r"^TableGroup (\w+)", dbml, flags=re.MULTILINE))
    decision_log = (
        DECISION_LOG_PATH.read_text(encoding="utf-8")
        if DECISION_LOG_PATH.exists()
        else ""
    )
    decisions = set(
        re.findall(r"^\| ([A-Z]{2}-S?\d{2}) \|", decision_log, re.MULTILINE)
    )
    future_items: set[str] = set()
    for path in FUTURE_PATHS:
        if path.exists():
            text = path.read_text(encoding="utf-8")
            future_items |= set(re.findall(r"^#{2,3} (\d+)\.", text, re.MULTILINE))
    return tables, groups, decisions, future_items


def _check_feature(
    feature: Feature,
    catalog: Catalog,
    references: tuple[set[str], set[str], set[str], set[str]],
) -> None:
    """Append every rule this feature breaks to ``catalog.problems``."""
    tables, groups, decisions, future_items = references
    problems = catalog.problems
    code = feature.code
    if feature.priority not in PRIORITIES:
        problems.append(f"{code}: priority must be one of {', '.join(PRIORITIES)}")
    if feature.release not in RELEASES:
        problems.append(f"{code}: release must be one of {', '.join(RELEASES)}")
    if feature.offer not in OFFER_LABELS:
        problems.append(f"{code}: offer must be one of {', '.join(OFFER_LABELS)}")
    for role in feature.users:
        if role not in catalog.roles:
            problems.append(f"{code}: unknown role {role!r}")
    for surface, status in feature.surfaces.items():
        if status not in STATUSES:
            problems.append(f"{code}: {surface} status must be one of {STATUSES}")
    for dependency in feature.depends_on:
        if dependency == code:
            problems.append(f"{code}: depends on itself")
        elif dependency not in catalog.features:
            problems.append(f"{code}: depends on unknown feature {dependency}")
    own_domain = catalog.domain_of(code).backend_domain
    backend_domains = groups | {
        domain.backend_domain for domain in catalog.domains if domain.backend_domain
    }
    for touched in feature.touches:
        if touched not in backend_domains:
            problems.append(f"{code}: touches unknown backend domain {touched!r}")
        elif touched == own_domain:
            problems.append(f"{code}: touches its own domain {touched!r}")
    for table in feature.related_tables:
        if table not in tables:
            problems.append(f"{code}: related table {table!r} is not in the DBML")
    for old_code in feature.old_codes:
        if not OLD_CODE_PATTERN.match(old_code):
            problems.append(f"{code}: old code {old_code!r} must look like F-A1")
    nfr_codes = {nfr.code for nfr in catalog.nfrs}
    for source in feature.sources:
        if not source.startswith(tuple(SOURCE_PREFIXES)):
            problems.append(f"{code}: source {source!r} has an unknown prefix")
        elif source.startswith("Decision ") and source[9:] not in decisions:
            problems.append(f"{code}: {source} is not in the decision log")
        elif source.startswith("future.md ") and source[10:] not in future_items:
            problems.append(f"{code}: {source} does not exist")
        elif source.startswith("NF-") and source not in nfr_codes:
            problems.append(f"{code}: {source} is not a non-functional requirement")
        elif source.startswith("PRD ") and not OLD_CODE_PATTERN.match(source[4:]):
            problems.append(f"{code}: {source} must name a PRD code like F-A1")


def _check_dependency_cycles(catalog: Catalog) -> None:
    """Append a problem for each dependency cycle (reported once per cycle)."""
    visiting: list[str] = []
    finished: set[str] = set()
    reported: set[frozenset[str]] = set()

    def visit(code: str) -> None:
        """Depth-first walk recording any path that returns to itself."""
        if code in finished or code not in catalog.features:
            return
        if code in visiting:
            cycle = visiting[visiting.index(code) :]
            if frozenset(cycle) not in reported:
                reported.add(frozenset(cycle))
                catalog.problems.append(
                    "dependency cycle: " + " → ".join([*cycle, code])
                )
            return
        visiting.append(code)
        for dependency in catalog.features[code].depends_on:
            # A self-dependency is already reported by _check_feature.
            if dependency != code:
                visit(dependency)
        visiting.pop()
        finished.add(code)

    for code in catalog.features:
        visit(code)


def load_catalog(source_path: Path) -> Catalog:
    """Parse the source and collect every rule it breaks.

    Raises:
        CatalogError: When the file is not YAML or lacks a top-level section.
    """
    try:
        raw = yaml.safe_load(source_path.read_text(encoding="utf-8"))
    except yaml.YAMLError as error:
        raise CatalogError(str(error)) from error
    if not isinstance(raw, dict) or not all(
        isinstance(raw.get(key), container)
        for key, container in (
            ("roles", dict),
            ("domains", list),
            ("features", list),
            ("nfrs", list),
        )
    ):
        raise CatalogError("expected top-level roles, domains, features and nfrs")
    problems: list[str] = []
    roles = {
        str(code): _text_pair(labels, f"role {code}", problems)
        for code, labels in raw["roles"].items()
    }
    domains: list[Domain] = []
    for raw_domain in raw["domains"]:
        missing = DOMAIN_KEYS - set(raw_domain)
        unknown = set(raw_domain) - DOMAIN_KEYS - DOMAIN_OPTIONAL_KEYS
        if missing or unknown:
            problems.append(
                f"domain {raw_domain.get('prefix')!r}: missing {sorted(missing)}, "
                f"unknown {sorted(unknown)}"
            )
            continue
        domains.append(
            Domain(
                **{key: " ".join(str(raw_domain[key]).split()) for key in DOMAIN_KEYS},
                backend_domain=raw_domain.get("backend_domain"),
                sheet_vi=raw_domain.get("sheet_vi"),
            )
        )
    for attribute in ("prefix", "key"):
        values = [getattr(domain, attribute) for domain in domains]
        for duplicate in sorted({value for value in values if values.count(value) > 1}):
            problems.append(f"domain {attribute} {duplicate!r} is used twice")
    prefixes = {domain.prefix for domain in domains}
    features: dict[str, Feature] = {}
    for raw_feature in raw["features"]:
        feature = _parse_feature(raw_feature, problems)
        if feature is None:
            continue
        if feature.code in features:
            problems.append(f"{feature.code}: code is used twice")
            continue
        if feature.prefix not in prefixes:
            problems.append(f"{feature.code}: no domain has prefix {feature.prefix}")
            continue
        features[feature.code] = feature
    nfrs: list[Nfr] = []
    for raw_nfr in raw["nfrs"]:
        code = str(raw_nfr.get("code"))
        if not NFR_CODE_PATTERN.match(code):
            problems.append(f"non-functional requirement {code!r}: code like NF-01")
            continue
        nfrs.append(
            Nfr(
                code=code,
                category=_text_pair(raw_nfr.get("category"), code, problems),
                requirement=_text_pair(raw_nfr.get("requirement"), code, problems),
                target=_text_pair(raw_nfr.get("target"), code, problems),
                features=_string_list(raw_nfr.get("features"), code, problems),
            )
        )
    catalog = Catalog(roles, domains, features, nfrs, problems)
    for domain in domains:
        domain.features = sorted(
            (
                feature
                for feature in features.values()
                if feature.prefix == domain.prefix
            ),
            key=lambda feature: feature.code,
        )
        if len(_sheet_name(domain)) > SHEET_NAME_LIMIT:
            problems.append(
                f"domain {domain.prefix}: sheet name {_sheet_name(domain)!r} is longer"
                f" than {SHEET_NAME_LIMIT} characters; add a shorter sheet_vi"
            )
        if not domain.features:
            problems.append(f"domain {domain.prefix} has no features")
    references = _read_reference_sets()
    for feature in features.values():
        _check_feature(feature, catalog, references)
    for nfr in nfrs:
        for code in nfr.features:
            if code not in features:
                problems.append(f"{nfr.code}: unknown feature {code}")
    _check_dependency_cycles(catalog)
    return catalog


# ---------------------------------------------------------------------------
# Markdown views
# ---------------------------------------------------------------------------


def _md_cell(text: str) -> str:
    """Make text safe inside a Markdown table cell."""
    return text.replace("|", "\\|").replace("\n", " ")


def _feature_link(code: str, catalog: Catalog, from_domain: Domain | None) -> str:
    """Markdown link to a feature's section, relative to the linking page.

    Args:
        code: Target feature code.
        catalog: The catalog.
        from_domain: Domain page the link sits on, or None for the README.
    """
    domain = catalog.domain_of(code)
    anchor = code.lower()
    if from_domain is None:
        return f"[{code}](domains/{domain.key}.md#{anchor})"
    if from_domain is domain:
        return f"[{code}](#{anchor})"
    return f"[{code}]({domain.key}.md#{anchor})"


def _surface_summary(feature: Feature) -> str:
    """Status of every surface, e.g. ``Backend ⬜ · App ⬜``."""
    return " · ".join(
        f"{SURFACE_LABELS[surface]} {STATUS_MARKS[status]}"
        for surface, status in feature.surfaces.items()
    )


def _progress(features: list[Feature], surface: str) -> str:
    """``done/needed`` for one surface, or ``—`` when no feature needs it."""
    needed = [feature for feature in features if surface in feature.surfaces]
    if not needed:
        return "—"
    done = sum(feature.surfaces[surface] == "done" for feature in needed)
    return f"{done}/{len(needed)}"


def render_readme(catalog: Catalog) -> str:
    """Render the catalog's front page."""
    all_features = list(catalog.features.values())
    lines = [
        GENERATED_HEADER.rstrip("\n"),
        "",
        "# Feature catalog",
        "",
        "> Every feature of the G3 Network product, for every user (driver app,",
        "> web portal for fleets, administrators and customer care, background",
        "> system), broken down into capabilities and tracked per surface",
        "> (backend, app, portal). Source: [`features.yaml`](features.yaml);",
        "> Vietnamese workbook for business readers: `features.xlsx`. Feature",
        "> codes are stable; the old `F-A1`-style codes are mapped at the end.",
        "",
        "## Progress",
        "",
        "A feature is done when every surface it needs is done. A cell shows",
        "*done / features that need this surface*.",
        "",
        "| Domain | Features | Done | Backend | App | Portal |",
        "|---|---|---|---|---|---|",
    ]
    for domain in catalog.domains:
        done_count = sum(feature.is_done for feature in domain.features)
        lines.append(
            f"| [{_md_cell(domain.title)}](domains/{domain.key}.md) "
            f"| {len(domain.features)} | {done_count} | "
            + " | ".join(_progress(domain.features, surface) for surface in SURFACES)
            + " |"
        )
    total_done = sum(feature.is_done for feature in all_features)
    lines.append(
        f"| **Total** | **{len(all_features)}** | **{total_done}** | "
        + " | ".join(f"**{_progress(all_features, s)}**" for s in SURFACES)
        + " |"
    )
    lines += [
        "",
        "## Domains and releases",
        "",
        "| Code | Domain | Backend domain | Scope | " + " | ".join(RELEASES) + " |",
        "|---|---|---|---|" + "---|" * len(RELEASES),
    ]
    for domain in catalog.domains:
        release_counts = [
            str(sum(feature.release == release for feature in domain.features) or "")
            for release in RELEASES
        ]
        backend = f"`{domain.backend_domain}`" if domain.backend_domain else "—"
        lines.append(
            f"| {domain.prefix} | [{_md_cell(domain.name)}](domains/{domain.key}.md) "
            f"| {backend} | {_md_cell(domain.summary)} | "
            + " | ".join(release_counts)
            + " |"
        )
    lines += [
        "",
        "## How to read a feature",
        "",
        "- **Status** per surface: "
        + ", ".join(f"{STATUS_MARKS[s]} {s}" for s in STATUSES)
        + ". A surface the feature does not need is not listed.",
        "- **Surfaces**: Backend (this repository), App (driver mobile / vehicle"
        " app), Portal (web portal: fleet, administration, customer-care console).",
        "- **Priority**: " + ", ".join(PRIORITIES) + " (MoSCoW).",
        "- **Release**: P1.0 launch, P1.1 about 90 days later, P1.5 2026–27, P2"
        " Phase 2.",
        "- **Offer**: "
        + "; ".join(f"*{label[0]}*" for label in OFFER_LABELS.values())
        + ". Prices come with the pricing work; internal-only features are"
        " never put in a plan.",
        "- **Related tables** stay empty until the database review is finished.",
        "- **Priority and release** follow the PRD for features it lists and"
        " the phase-1 data sheet for features it does not; where both name a"
        " release, the data sheet (the newer document) wins.",
        "",
        "### Roles",
        "",
        "| Code | Role | Vai trò |",
        "|---|---|---|",
    ]
    lines += [
        f"| `{code}` | {english} | {vietnamese} |"
        for code, (english, vietnamese) in catalog.roles.items()
    ]
    lines += [
        "",
        "## Non-functional requirements",
        "",
        "| Code | Category | Requirement | Target | Features |",
        "|---|---|---|---|---|",
    ]
    for nfr in catalog.nfrs:
        feature_links = ", ".join(
            _feature_link(code, catalog, None) for code in nfr.features
        )
        lines.append(
            f"| {nfr.code} | {_md_cell(nfr.category[0])} "
            f"| {_md_cell(nfr.requirement[0])} | {_md_cell(nfr.target[0])} "
            f"| {feature_links or 'platform-wide'} |"
        )
    old_map: dict[str, list[str]] = {}
    for feature in catalog.features.values():
        for old_code in feature.old_codes:
            old_map.setdefault(old_code, []).append(feature.code)
    lines += [
        "",
        "## Old feature codes",
        "",
        "Where each code of the old `feature-list.md` went.",
        "",
        "| Old code | Features |",
        "|---|---|",
    ]
    for old_code in sorted(old_map, key=lambda c: (c[2], int(c[3:]))):
        links = ", ".join(
            _feature_link(code, catalog, None) for code in sorted(old_map[old_code])
        )
        lines.append(f"| {old_code} | {links} |")
    return "\n".join(lines) + "\n"


def _render_feature(feature: Feature, catalog: Catalog, domain: Domain) -> list[str]:
    """Render one feature's detail section on its domain page."""
    roles = ", ".join(
        catalog.roles[role][0] for role in feature.users if role in catalog.roles
    )
    lines = [
        f'<a id="{feature.code.lower()}"></a>',
        "",
        f"### {feature.code} {feature.name}",
        "",
        f"*{feature.name_vi}* · {feature.priority} · {feature.release} · "
        f"{OFFER_LABELS.get(feature.offer, (feature.offer,))[0]}",
        "",
        feature.what,
        "",
        f"**Value:** {feature.value}",
        "",
        f"**Users:** {roles}",
        "",
        "**Capabilities:**",
        "",
    ]
    lines += [f"- {english}" for english, _ in feature.capabilities]
    lines += ["", f"**Status:** {_surface_summary(feature)}", ""]
    details = []
    if feature.depends_on:
        details.append(
            "**Depends on:** "
            + ", ".join(_feature_link(c, catalog, domain) for c in feature.depends_on)
        )
    dependents = catalog.dependents_of(feature.code)
    if dependents:
        details.append(
            "**Needed by:** "
            + ", ".join(_feature_link(c, catalog, domain) for c in dependents)
        )
    if feature.touches:
        details.append(
            "**Also touches:** " + ", ".join(f"`{name}`" for name in feature.touches)
        )
    details.append(
        "**Related tables:** "
        + (
            ", ".join(f"`{table}`" for table in feature.related_tables)
            or "— (after the database review)"
        )
    )
    details.append("**Sources:** " + " · ".join(feature.sources))
    if feature.old_codes:
        details.append("**Old codes:** " + ", ".join(feature.old_codes))
    lines += [detail + "  " for detail in details[:-1]] + [details[-1], ""]
    if feature.questions:
        lines += ["**Open questions:**", ""]
        lines += [f"- {english}" for english, _ in feature.questions]
        lines.append("")
    return lines


def render_domain_page(domain: Domain, catalog: Catalog) -> str:
    """Render one domain's checklist and feature details."""
    backend = (
        f"Backend domain: `{domain.backend_domain}`."
        if domain.backend_domain
        else "No backend domain yet."
    )
    lines = [
        GENERATED_HEADER.rstrip("\n"),
        "",
        f"# {domain.title}",
        "",
        f"*{domain.name_vi}* · [← Feature catalog](../README.md)",
        "",
        f"{domain.summary} {backend}",
        "",
        "## Checklist",
        "",
    ]
    for feature in domain.features:
        box = "x" if feature.is_done else " "
        lines.append(
            f"- [{box}] **{feature.code}** [{feature.name}](#{feature.code.lower()})"
            f" — {_surface_summary(feature)}"
        )
    lines += ["", "## Features", ""]
    for feature in domain.features:
        lines += _render_feature(feature, catalog, domain)
    return "\n".join(lines).rstrip("\n") + "\n"


def render_all(catalog: Catalog) -> dict[Path, str]:
    """Render every Markdown view, keyed by output path."""
    rendered = {README_PATH: render_readme(catalog)}
    for domain in catalog.domains:
        rendered[DOMAINS_DIR / f"{domain.key}.md"] = render_domain_page(domain, catalog)
    return rendered


def _stale_generated_files(rendered: dict[Path, str]) -> list[Path]:
    """Return generated domain pages whose domain no longer exists."""
    if not DOMAINS_DIR.exists():
        return []
    return [
        path
        for path in DOMAINS_DIR.glob("*.md")
        if path not in rendered
        and path.read_text(encoding="utf-8").startswith(GENERATED_MARKER)
    ]


# ---------------------------------------------------------------------------
# Excel workbook (Vietnamese)
# ---------------------------------------------------------------------------

# The workbook is for business readers (BOD, BA), so every label and text in
# it is Vietnamese: labels come from the constants below, content from the
# source's *_vi fields and pairs.
OVERVIEW_SHEET = "Tổng quan"
NFR_SHEET = "Phi chức năng"
SHEET_NAME_LIMIT = 31
XLSX_FONT = "Arial"
XLSX_HEADER_FILL = "1F4E78"
XLSX_DOMAIN_FILL = "D9E1F2"
XLSX_LINK_COLOR = "0563C1"
XLSX_MUTED_COLOR = "595959"
XLSX_OPEN_LINK_LABEL = "→ Mở"
XLSX_BULLET = "• "
OVERVIEW_HEADERS = (
    "STT",
    "Mã",
    "Nhóm tính năng",
    "Phạm vi",
    "Domain backend",
    "Số tính năng",
    *RELEASES,
    "Backend xong",
    "App xong",
    "Portal xong",
    "Mở",
)
OVERVIEW_WIDTHS = [5, 7, 32, 60, 18, 10, 7, 7, 7, 7, 11, 10, 11, 8]
DOMAIN_HEADERS = (
    "Mã",
    "Tên tính năng",
    "Mô tả",
    "Giá trị mang lại",
    "Người dùng",
    "Chức năng chi tiết",
    "Ưu tiên",
    "Đợt",
    "Hình thức bán / Giá",
    *(VI_SURFACE_LABELS[surface] for surface in SURFACES),
    "Phụ thuộc",
    "Bảng dữ liệu liên quan",
    "Câu hỏi mở",
    "Nguồn",
    "Mã cũ",
)
DOMAIN_WIDTHS = [8, 26, 48, 36, 22, 56, 10, 7, 16, 10, 10, 10, 16, 18, 44, 26, 9]
NFR_HEADERS = ("Mã", "Nhóm", "Yêu cầu", "Chỉ tiêu", "Tính năng liên quan")
NFR_WIDTHS = [8, 24, 34, 70, 30]


class _SheetWriter:
    """Row-by-row writer for one worksheet with this workbook's styling.

    Attributes:
        worksheet: The openpyxl worksheet being written.
        row: Next row to write (1-based).
        widths: Column widths in characters, used to estimate wrapped heights.
    """

    def __init__(self, worksheet: object, widths: list[int]) -> None:
        """Set column widths and start at row 1."""
        from openpyxl.utils import get_column_letter  # noqa: PLC0415

        self.worksheet = worksheet
        self.row = 1
        self.widths = widths
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width

    def write(
        self,
        values: list[object],
        *,
        bold: bool = False,
        size: int = 10,
        color: str | None = None,
        fill: str | None = None,
        links: dict[int, str] | None = None,
    ) -> int:
        """Write one wrapped row and return its number.

        Args:
            values: Cell values from column A; None or "" leaves a cell empty.
            bold: Bold font for the whole row.
            size: Font size in points.
            color: Font color (hex RGB) for the whole row.
            fill: Background fill (hex RGB) for every written cell.
            links: Internal hyperlinks ``{1-based column: sheet name}``.
        """
        from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
        from openpyxl.worksheet.hyperlink import Hyperlink  # noqa: PLC0415

        line_count = 1
        for column_index, value in enumerate(values, start=1):
            # An empty string is saved as an empty cell and reads back as
            # None; write None up front so a round trip compares equal.
            if value == "":
                value = None
            cell = self.worksheet.cell(row=self.row, column=column_index, value=value)
            link_target = (links or {}).get(column_index)
            cell.font = Font(
                name=XLSX_FONT,
                size=size,
                bold=bold,
                color=XLSX_LINK_COLOR if link_target else color,
                underline="single" if link_target else None,
            )
            if link_target:
                # An internal "location" link, so Excel and LibreOffice jump
                # to the sheet in place.
                cell.hyperlink = Hyperlink(
                    ref=cell.coordinate, location=f"'{link_target}'!A1"
                )
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if value is not None and column_index <= len(self.widths):
                # openpyxl cannot auto-fit and Excel does not re-fit rows on
                # open, so estimate the wrapped line count from the width.
                chars_per_line = max(self.widths[column_index - 1] - 2, 1)
                cell_lines = sum(
                    max(1, -(-len(part) // chars_per_line))
                    for part in str(value).split("\n")
                )
                line_count = max(line_count, cell_lines)
        if line_count > 1:
            self.worksheet.row_dimensions[self.row].height = 13.5 * line_count
        self.row += 1
        return self.row - 1

    def header(self, labels: list[str]) -> int:
        """Write a white-on-blue header row and return its number."""
        return self.write(labels, bold=True, color="FFFFFF", fill=XLSX_HEADER_FILL)

    def blank(self) -> None:
        """Skip one row."""
        self.row += 1


def _sheet_name(domain: Domain) -> str:
    """Sheet name of a domain: prefix and Vietnamese name (or ``sheet_vi``).

    ``load_catalog`` reports a name longer than Excel's limit, so it is never
    cut in the middle of a word.
    """
    name = f"{domain.prefix} {domain.sheet_vi or domain.name_vi}"
    for character in "[]:*?/\\":
        name = name.replace(character, "-")
    return name


def _vi_source(source: str) -> str:
    """Source reference with its prefix in Vietnamese."""
    for prefix, vi_prefix in SOURCE_PREFIXES.items():
        if source.startswith(prefix):
            return vi_prefix + source[len(prefix) :]
    return source


def _bullets(items: list[str]) -> str:
    """Join items into one cell, one bullet per line."""
    return "\n".join(XLSX_BULLET + item for item in items)


def _write_overview_sheet(worksheet: object, catalog: Catalog) -> None:
    """Fill the overview sheet: domain index first, then the legend."""
    writer = _SheetWriter(worksheet, OVERVIEW_WIDTHS)
    all_features = list(catalog.features.values())
    writer.write(["DANH MỤC TÍNH NĂNG — G3 NETWORK"], bold=True, size=14)
    header_row = writer.header(list(OVERVIEW_HEADERS))
    # Freeze only the title and header: the index is the first thing on the
    # sheet, and freezing more would leave no room to scroll.
    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
    for index, domain in enumerate(catalog.domains, start=1):
        releases = [
            sum(feature.release == release for feature in domain.features) or None
            for release in RELEASES
        ]
        writer.write(
            [
                index,
                domain.prefix,
                domain.name_vi,
                domain.summary_vi,
                domain.backend_domain or "Chưa có",
                len(domain.features),
                *releases,
                *(_progress(domain.features, surface) for surface in SURFACES),
                XLSX_OPEN_LINK_LABEL,
            ],
            links={3: _sheet_name(domain), len(OVERVIEW_HEADERS): _sheet_name(domain)},
        )
    writer.write(
        [
            None,
            None,
            "Tổng cộng",
            None,
            None,
            len(all_features),
            *(
                sum(feature.release == release for feature in all_features)
                for release in RELEASES
            ),
            *(_progress(all_features, surface) for surface in SURFACES),
        ],
        bold=True,
        fill=XLSX_DOMAIN_FILL,
    )
    writer.write(
        [None, None, "Yêu cầu phi chức năng", None, None, len(catalog.nfrs)],
        links={3: NFR_SHEET},
    )
    writer.blank()
    writer.write([None, "CÁCH ĐỌC"], bold=True, size=12)
    legend = [
        (
            "Trạng thái",
            "Theo từng phần cần làm (Backend, App, Portal web): "
            + ", ".join(VI_STATUS_LABELS[s] for s in STATUSES)
            + f". '{VI_NOT_NEEDED}' = tính năng không cần phần đó. Cột 'xong' ở trên"
            " = số tính năng đã xong / số tính năng cần phần đó.",
        ),
        ("Ưu tiên", ", ".join(f"{VI_PRIORITY_LABELS[p]} ({p})" for p in PRIORITIES)),
        ("Đợt", "P1.0 ra mắt, P1.1 khoảng 90 ngày sau, P1.5 năm 2026–27, P2 Phase 2."),
        (
            "Hình thức bán",
            "; ".join(label[1] for label in OFFER_LABELS.values())
            + ". Giá được bổ sung khi làm phần định giá; tính năng nội bộ không bao"
            " giờ nằm trong gói.",
        ),
        ("Bảng dữ liệu", "Để trống cho tới khi rà soát xong thiết kế cơ sở dữ liệu."),
    ]
    for label, text in legend:
        writer.write([None, None, label, text])
    writer.blank()
    writer.write([None, "VAI TRÒ"], bold=True, size=12)
    for code, (_, vietnamese) in catalog.roles.items():
        writer.write([None, None, vietnamese, code], color=None)


def _write_domain_sheet(worksheet: object, domain: Domain, catalog: Catalog) -> None:
    """Fill one domain's sheet: title, scope, then one row per feature."""
    writer = _SheetWriter(worksheet, DOMAIN_WIDTHS)
    writer.write([f"{domain.prefix} — {domain.name_vi}"], bold=True, size=14)
    writer.write([domain.summary_vi], color=XLSX_MUTED_COLOR)
    writer.write([f"← {OVERVIEW_SHEET}"], links={1: OVERVIEW_SHEET})
    header_row = writer.header(list(DOMAIN_HEADERS))
    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=3)
    for feature in domain.features:
        writer.write(
            [
                feature.code,
                feature.name_vi,
                feature.what_vi,
                feature.value_vi,
                ", ".join(
                    catalog.roles[r][1] for r in feature.users if r in catalog.roles
                ),
                _bullets([vietnamese for _, vietnamese in feature.capabilities]),
                VI_PRIORITY_LABELS.get(feature.priority, feature.priority),
                feature.release,
                OFFER_LABELS.get(feature.offer, ("", feature.offer))[1],
                *(
                    VI_STATUS_LABELS.get(feature.surfaces[surface], "")
                    if surface in feature.surfaces
                    else VI_NOT_NEEDED
                    for surface in SURFACES
                ),
                ", ".join(feature.depends_on),
                ", ".join(feature.related_tables),
                _bullets([vietnamese for _, vietnamese in feature.questions]),
                ", ".join(_vi_source(source) for source in feature.sources),
                ", ".join(feature.old_codes),
            ]
        )


def _write_nfr_sheet(worksheet: object, catalog: Catalog) -> None:
    """Fill the non-functional requirements sheet."""
    writer = _SheetWriter(worksheet, NFR_WIDTHS)
    writer.write(["YÊU CẦU PHI CHỨC NĂNG"], bold=True, size=14)
    writer.write([f"← {OVERVIEW_SHEET}"], links={1: OVERVIEW_SHEET})
    header_row = writer.header(list(NFR_HEADERS))
    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
    for nfr in catalog.nfrs:
        writer.write(
            [
                nfr.code,
                nfr.category[1],
                nfr.requirement[1],
                nfr.target[1],
                ", ".join(nfr.features) or "Toàn nền tảng",
            ]
        )


def build_workbook(catalog: Catalog) -> object:
    """Build the Vietnamese workbook: overview, one sheet per domain, NFRs."""
    from openpyxl import Workbook  # noqa: PLC0415

    workbook = Workbook()
    overview_sheet = workbook.active
    overview_sheet.title = OVERVIEW_SHEET
    _write_overview_sheet(overview_sheet, catalog)
    for domain in catalog.domains:
        _write_domain_sheet(workbook.create_sheet(_sheet_name(domain)), domain, catalog)
    _write_nfr_sheet(workbook.create_sheet(NFR_SHEET), catalog)
    return workbook


def _workbook_values(workbook: object) -> list[tuple[str, list[tuple[object, ...]]]]:
    """Every sheet's title and cell values, for comparing two workbooks.

    Byte comparison is useless for .xlsx (the zip holds timestamps), so
    ``generate`` and ``check`` compare values instead.
    """
    return [
        (sheet.title, [tuple(row) for row in sheet.iter_rows(values_only=True)])
        for sheet in workbook.worksheets
    ]


def _is_workbook_current(workbook: object) -> bool:
    """Whether the committed .xlsx holds the same values as ``workbook``."""
    from openpyxl import load_workbook  # noqa: PLC0415

    if not XLSX_PATH.exists():
        return False
    return _workbook_values(load_workbook(XLSX_PATH)) == _workbook_values(workbook)


# ---------------------------------------------------------------------------
# Commands
# ---------------------------------------------------------------------------


def command_generate(catalog: Catalog) -> int:
    """Write every view; refuse while the source breaks a rule."""
    if catalog.problems:
        for problem in catalog.problems:
            print(f"PROBLEM {problem}")
        print(f"{len(catalog.problems)} problem(s); nothing written")
        return 1
    rendered = render_all(catalog)
    DOMAINS_DIR.mkdir(parents=True, exist_ok=True)
    for path, content in rendered.items():
        path.write_text(content, encoding="utf-8")
        print(f"wrote {path.relative_to(REPO_ROOT)}")
    for path in _stale_generated_files(rendered):
        path.unlink()
        print(f"removed {path.relative_to(REPO_ROOT)}")
    # Rewrite the workbook only when its content changed: every save embeds
    # a fresh timestamp, which would otherwise make git see a new binary.
    workbook = build_workbook(catalog)
    if _is_workbook_current(workbook):
        print(f"unchanged {XLSX_PATH.relative_to(REPO_ROOT)}")
    else:
        workbook.save(XLSX_PATH)
        print(f"wrote {XLSX_PATH.relative_to(REPO_ROOT)}")
    return 0


def command_check(catalog: Catalog) -> int:
    """Check the source's rules and that every view is current."""
    problems = list(catalog.problems)
    if not problems:
        rendered = render_all(catalog)
        for path, content in rendered.items():
            if not path.exists() or path.read_text(encoding="utf-8") != content:
                problems.append(f"{path.relative_to(REPO_ROOT)}: stale, run generate")
        for path in _stale_generated_files(rendered):
            problems.append(
                f"{path.relative_to(REPO_ROOT)}: domain removed, run generate"
            )
        if not _is_workbook_current(build_workbook(catalog)):
            problems.append(f"{XLSX_PATH.relative_to(REPO_ROOT)}: stale, run generate")
    for problem in problems:
        print(f"PROBLEM {problem}")
    if problems:
        print(f"{len(problems)} problem(s) found")
        return 1
    print(
        f"OK: {len(catalog.features)} features in {len(catalog.domains)} domains; "
        "views are up to date"
    )
    return 0


def main() -> int:
    """Parse arguments, load the catalog, and run the chosen command."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("command", choices=("generate", "check"))
    arguments = parser.parse_args()
    try:
        catalog = load_catalog(SOURCE_PATH)
    except CatalogError as error:
        print(f"ERROR in {SOURCE_PATH.relative_to(REPO_ROOT)}:\n{error}")
        return 2
    if arguments.command == "generate":
        return command_generate(catalog)
    return command_check(catalog)


if __name__ == "__main__":
    sys.exit(main())

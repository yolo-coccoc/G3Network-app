"""Generate the domain-model views from the single DBML source, and check it.

The whole database design lives in one file,
``docs/design/domain-model/domain-model.dbml``. This script renders
three levels of view from it and verifies it against the backend:

- ``generate``: writes ``overview.md`` (L1: domain map, ownership, open
  decisions, for the BOD), ``domains/<domain>.md`` (L2: one entity
  diagram per domain; L3: column-level tables, indexes, references), and
  ``domain-model.xlsx`` (in Vietnamese, for business readers: a "Tổng
  quan" sheet indexing every domain and table with links, then one sheet
  per table with each column's meaning and example; no build progress).
- ``check``: fails when a ``@status built`` table no longer matches the
  SQLAlchemy models (tables, columns, types, nullability, primary keys,
  foreign keys, enum values), when a column lacks a meaning or an example,
  or when any generated view (Markdown or Excel) is stale.

Run from the repository root (``check`` needs the backend environment to
import the models; ``generate`` only needs pydbml and openpyxl)::

    uv run --project backend --with pydbml --with openpyxl python \\
        .claude/skills/domain-model/scripts/domain_model.py generate
    uv run --project backend --with pydbml --with openpyxl python \\
        .claude/skills/domain-model/scripts/domain_model.py check

Limitations: ``check`` does not compare indexes, unique constraints,
defaults, or check constraints; those are documented in the DBML but only
reviewed by hand.
"""

from __future__ import annotations

import argparse
import importlib
import pkgutil
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from pydbml import PyDBML

# The script lives at <root>/.claude/skills/domain-model/scripts/, so the
# repository root is four levels up.
REPO_ROOT = Path(__file__).resolve().parents[4]
MODEL_DIR = REPO_ROOT / "docs" / "design" / "domain-model"
SOURCE_PATH = MODEL_DIR / "domain-model.dbml"
OVERVIEW_PATH = MODEL_DIR / "overview.md"
DOMAINS_DIR = MODEL_DIR / "domains"

GENERATED_MARKER = "<!-- GENERATED from domain-model.dbml"
GENERATED_HEADER = (
    f"{GENERATED_MARKER} by .claude/skills/domain-model/scripts/domain_model.py"
    " - do not edit by hand; edit the .dbml and regenerate. -->\n"
)

STATUSES = ("built", "planned", "proposed")
OWNERS = ("customer", "internal", "two-party", "undecided")
STATUS_BADGES = {
    "built": "✅ built",
    "planned": "📋 planned",
    "proposed": "🆕 proposed",
}
OWNER_MEANINGS = {
    "customer": "Belongs to one organization (normally a customer; a G3 company can own such rows too); carries `organization_id` or reads it through its parent (DM-24).",
    "internal": "Our own data (the organization running the platform), shared across all organizations.",
    "two-party": "A G3 asset used by a customer (both have a stake).",
    "undecided": "Waiting on an open decision.",
}
TENANT_COLUMN = "organization_id"
# The tenant table. Links to it (one per customer-owned table) are left out
# of the L1 domain map, which would otherwise be dominated by them, and are
# summarized in one sentence instead.
TENANT_TABLE = "organizations"
PLANNED_COLUMN_TAG = "@planned"
# Target design of a built table (decision PR-11: design first, refactor in
# bulk): a built column, or a value of a built enum, that the refactor drops.
# It still exists in the code, so `check` compares it as usual; the views mark
# it and generated history tables leave it out.
REMOVED_TAG = "@remove"
# Every column note ends with "@example <sample value>"; the text before it
# is the column's meaning. Both are required (check fails without them).
EXAMPLE_TAG = "@example"
# Vietnamese text for the Excel export (the Markdown views stay English).
# Column notes carry it inline ("English. @vi Tiếng Việt. @example v");
# table, domain and free-standing notes carry a "@vi-name <tên>" line (tables
# and domains only) and a "@vi <text...>" block that runs to the note's end.
VI_TAG = "@vi"
VI_NAME_TAG = "@vi-name"
# Change history (owner decision D8, .claude/rules/database.md): a table tags
# the columns whose change is audited ("@tracked a,b,c"); its history table
# ("@history-of <table>") mirrors every source column and adds only these.
HISTORY_METADATA_COLUMNS = ("history_id", "changed_at", "changed_by", "change_reason")
# "@tracked *" = history covers the whole table (minus "@untracked" columns).
TRACK_ALL = "*"
# Companion tables of a main table: generated history (numbered <N>.h) and
# the hand-designed "@state-of" live-state table (numbered like any table).
COMPANION_LABELS = {"history": "change history", "state": "live state"}
VI_COMPANION_LABELS = {"history": "Bảng lịch sử của", "state": "Bảng trạng thái của"}
TRACKED_MARK = "🔍"
TAG_PATTERN = re.compile(r"@([\w-]+)((?:\s+[^@\s]+)*)")

# Mermaid classes for the L1 domain map. Mid-tone fills with dark text stay
# legible on both light and dark GitHub/VS Code themes.
DOMAIN_MAP_CLASSES = (
    "  classDef built fill:#C8E6C9,stroke:#2E7D32,color:#1B5E20\n"
    "  classDef partial fill:#FFE0B2,stroke:#EF6C00,color:#6D3100\n"
    "  classDef planned fill:#ECEFF1,stroke:#78909C,color:#37474F,"
    "stroke-dasharray:5 5\n"
)


class DomainModelError(Exception):
    """Raised when the DBML source breaks one of the model conventions."""


@dataclass(frozen=True)
class ColumnInfo:
    """One column of a table, as written in the DBML source.

    Attributes:
        name: Column name.
        type_label: Type as written (enum name for enum columns).
        enum_values: Enum member names when the type is a DBML Enum, else None.
        is_pk: Part of the primary key.
        is_not_null: Declared ``not null`` (primary-key columns count as not null).
        is_unique: Declared ``unique`` on the column itself.
        meaning: What the column holds, in English (note text without the tags).
        vi_meaning: The same meaning in Vietnamese (``@vi``), may be empty.
        example: Sample value from the ``@example`` tag, may be empty.
        is_planned: Proposed column on a built table (note starts with ``@planned``).
        planned_detail: Text after ``@planned`` (e.g. decision IDs), may be empty.
        is_removed: Built column the target design drops (note starts with
            ``@remove``).
        removed_detail: Text after ``@remove`` (e.g. decision IDs), may be empty.
        removed_enum_values: Enum values whose note starts with ``@remove``.
        planned_enum_values: Enum values whose note starts with ``@planned``:
            proposed for a built enum, not in the code yet.
    """

    name: str
    type_label: str
    enum_values: tuple[str, ...] | None
    is_pk: bool
    is_not_null: bool
    is_unique: bool
    meaning: str
    vi_meaning: str
    example: str
    is_planned: bool
    planned_detail: str
    is_removed: bool = False
    removed_detail: str = ""
    removed_enum_values: tuple[str, ...] = ()
    planned_enum_values: tuple[str, ...] = ()


@dataclass(frozen=True)
class IndexInfo:
    """One index declared in a table's ``indexes`` block.

    Attributes:
        label: Human-readable column/expression list.
        name: Index name, may be empty.
        is_unique: Declared unique.
        note: Index note (partial predicate, index type ...), may be empty.
    """

    label: str
    name: str
    is_unique: bool
    note: str


@dataclass(frozen=True)
class RefInfo:
    """One foreign key, normalized so ``from_*`` is always the FK holder.

    Attributes:
        from_table: Table holding the foreign-key column.
        from_column: Foreign-key column.
        to_table: Referenced table.
        to_column: Referenced column.
        kind: ``many-to-one``, ``one-to-one`` or ``many-to-many``.
        on_delete: Delete action as written (lower case), or empty.
        is_planned: True unless both tables are built and the column is not
            a ``@planned`` column.
    """

    from_table: str
    from_column: str
    to_table: str
    to_column: str
    kind: str
    on_delete: str
    is_planned: bool


@dataclass
class TableInfo:
    """One table with its parsed tags.

    Attributes:
        name: Table name.
        domain: Owning TableGroup (domain) name.
        status: ``built``, ``planned`` or ``proposed``.
        owner: Data owner, one of ``OWNERS``.
        features: Feature codes from the ``@features`` tag.
        hypertable_column: Time column when tagged ``@hypertable``, else empty.
        description: English table note without the tag line or ``@vi`` part.
        vi_name: Vietnamese table name (``@vi-name``), may be empty.
        vi_description: Vietnamese description (``@vi``), may be empty.
        columns: Columns in declaration order.
        indexes: Declared indexes.
        tracked_columns: Columns whose change copies the whole old row into
            this table's history table: every non-PK column for
            ``@tracked *`` minus ``@untracked`` ones, or an explicit list.
        kind: ``main``, ``history`` (generated by the tool, never written in
            the DBML) or ``state`` (a hand-designed ``@state-of`` companion).
        parent: For a history or state table, its main table; empty otherwise.
        history_table: For a tracked table, the name of its generated
            history table; empty otherwise.
        state_table: For a main table with a state companion, its name.
        number: Display number: ``1``, ``2``... shared by main and state
            tables in design order; ``<N>.h`` for a generated history table.
    """

    name: str
    domain: str
    status: str
    owner: str
    features: list[str]
    hypertable_column: str
    description: str
    vi_name: str
    vi_description: str
    columns: list[ColumnInfo]
    indexes: list[IndexInfo]
    tracked_columns: list[str] = field(default_factory=list)
    kind: str = "main"
    parent: str = ""
    history_table: str = ""
    state_table: str = ""
    number: str = ""

    def column(self, column_name: str) -> ColumnInfo | None:
        """Return the column with this name, or None."""
        for column_info in self.columns:
            if column_info.name == column_name:
                return column_info
        return None


@dataclass
class DomainInfo:
    """One domain (DBML TableGroup).

    Attributes:
        name: Group name, which matches the backend domain directory.
        note: English group note (description and relationship sentences, Markdown).
        vi_title: Vietnamese domain title (``@vi-name``), may be empty.
        vi_note: Vietnamese group note (``@vi``), may be empty.
        table_names: Table names in group order.
    """

    name: str
    note: str
    vi_title: str
    vi_note: str
    table_names: list[str]

    @property
    def title(self) -> str:
        """Human-readable domain title, e.g. ``Charging sessions``."""
        return self.name.replace("_", " ").capitalize()


@dataclass
class DomainModel:
    """The whole parsed model.

    Attributes:
        project_note: Project-level Markdown note.
        domains: Domains in source order.
        tables: Tables by name.
        refs: Normalized foreign keys in source order.
        sticky_notes: (name, English Markdown, Vietnamese Markdown) of
            free-standing notes.
        warnings: Non-fatal convention findings.
        documentation_gaps: Missing meanings, examples or Vietnamese text;
            ``generate`` reports them as warnings, ``check`` fails on them.
    """

    project_note: str
    domains: list[DomainInfo]
    tables: dict[str, TableInfo]
    refs: list[RefInfo]
    sticky_notes: list[tuple[str, str, str]]
    warnings: list[str] = field(default_factory=list)
    documentation_gaps: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------


def _note_text(note: object) -> str:
    """Return the text of a pydbml note object, or an empty string."""
    text = getattr(note, "text", None) if note is not None else None
    return (text or "").strip()


def _parse_tags(tag_line: str) -> dict[str, str]:
    """Parse ``@key value @key2 value2`` into a dict of stripped strings."""
    return {
        match.group(1): match.group(2).strip()
        for match in TAG_PATTERN.finditer(tag_line)
    }


def _split_vietnamese(text: str) -> tuple[str, str, str]:
    """Split a multi-line note into its English part and its Vietnamese parts.

    Args:
        text: Note text that may contain a ``@vi-name <name>`` line and a
            ``@vi <text>`` block running to the end of the note.

    Returns:
        ``(english_text, vi_name, vi_text)``; missing parts are empty strings.
    """
    english_lines: list[str] = []
    vi_lines: list[str] = []
    vi_name = ""
    is_in_vi_block = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith(f"{VI_NAME_TAG} "):
            vi_name = stripped[len(VI_NAME_TAG) :].strip()
        elif stripped == VI_TAG or stripped.startswith(f"{VI_TAG} "):
            is_in_vi_block = True
            vi_lines.append(stripped[len(VI_TAG) :].strip())
        elif is_in_vi_block:
            vi_lines.append(line)
        else:
            english_lines.append(line)
    return "\n".join(english_lines).strip(), vi_name, "\n".join(vi_lines).strip()


def _split_leading_tag(note: str, tag: str) -> tuple[bool, str, str]:
    """Split a leading ``@planned``/``@remove`` tag off a note.

    Accepted forms: "@tag", "@tag D1 D3", "@tag D1: text", "@tag: text".
    Only text before a colon counts as the detail.

    Returns:
        ``(has the tag, detail, remaining note)``.
    """
    if not note.startswith(tag):
        return False, "", note
    remainder = note[len(tag) :].strip()
    if ":" in remainder:
        detail, _, rest = remainder.partition(":")
    else:
        detail, rest = remainder, ""
    return True, detail.strip(), rest.strip()


def _parse_column(pydbml_column: object) -> ColumnInfo:
    """Convert a pydbml column into a ``ColumnInfo``."""
    column_type = pydbml_column.type
    enum_values: tuple[str, ...] | None = None
    removed_enum_values: tuple[str, ...] = ()
    planned_enum_values: tuple[str, ...] = ()
    if hasattr(column_type, "items"):
        type_label = column_type.name
        enum_values = tuple(item.name for item in column_type.items)
        removed_enum_values = tuple(
            item.name
            for item in column_type.items
            if _note_text(item.note).startswith(REMOVED_TAG)
        )
        planned_enum_values = tuple(
            item.name
            for item in column_type.items
            if _note_text(item.note).startswith(PLANNED_COLUMN_TAG)
        )
    else:
        type_label = str(column_type)
    note, _, example = _note_text(pydbml_column.note).partition(EXAMPLE_TAG)
    note, _, vi_meaning = note.partition(f" {VI_TAG} ")
    note, vi_meaning, example = note.strip(), vi_meaning.strip(), example.strip()
    is_planned, planned_detail, note = _split_leading_tag(note, PLANNED_COLUMN_TAG)
    is_removed, removed_detail, note = _split_leading_tag(note, REMOVED_TAG)
    return ColumnInfo(
        name=pydbml_column.name,
        type_label=type_label,
        enum_values=enum_values,
        is_pk=bool(pydbml_column.pk),
        is_not_null=bool(pydbml_column.not_null or pydbml_column.pk),
        is_unique=bool(pydbml_column.unique),
        meaning=note,
        vi_meaning=vi_meaning,
        example=example,
        is_planned=is_planned,
        planned_detail=planned_detail,
        is_removed=is_removed,
        removed_detail=removed_detail,
        removed_enum_values=removed_enum_values,
        planned_enum_values=planned_enum_values,
    )


def _parse_index(pydbml_index: object) -> IndexInfo:
    """Convert a pydbml index into an ``IndexInfo``."""
    # A column subject has a name; an expression subject (`lower(email)`)
    # carries its SQL in `text`. Never trim parentheses from it.
    subject_labels = [
        getattr(subject, "name", None) or getattr(subject, "text", str(subject))
        for subject in pydbml_index.subjects
    ]
    return IndexInfo(
        label=", ".join(subject_labels),
        name=pydbml_index.name or "",
        is_unique=bool(pydbml_index.unique),
        note=_note_text(pydbml_index.note),
    )


def _resolve_tracked_columns(
    table_name: str, tags: dict[str, str], columns: list[ColumnInfo]
) -> tuple[list[str], list[str]]:
    """Expand a table's ``@tracked``/``@untracked`` tags into column names.

    The change-history rule tracks a whole table: ``@tracked *`` means every
    column except the primary key (which never changes) and the columns
    listed in ``@untracked`` - kept only for machine-updated columns that
    would flood the history (e.g. a heartbeat time). An explicit
    ``@tracked a,b`` list is still accepted.

    Returns:
        ``(tracked column names, convention errors)``.
    """
    tracked_spec = [
        name for name in re.split(r"[,\s]+", tags.get("tracked", "")) if name
    ]
    untracked = [
        name for name in re.split(r"[,\s]+", tags.get("untracked", "")) if name
    ]
    column_names = [column.name for column in columns]
    errors = [
        f"{table_name}: @untracked names unknown column {name}"
        for name in untracked
        if name not in column_names
    ]
    if untracked and tracked_spec != [TRACK_ALL]:
        errors.append(f"{table_name}: @untracked needs @tracked {TRACK_ALL}")
    if tracked_spec == [TRACK_ALL]:
        tracked = [
            column.name
            for column in columns
            if not column.is_pk
            and not column.is_removed
            and column.name not in untracked
        ]
        return tracked, errors
    return tracked_spec, errors


def _singular(table_name: str) -> str:
    """Singular form of a plural snake_case table name (``fleets`` -> ``fleet``).

    Only the last word changes; covers the plural forms used in this model
    (``-ies`` -> ``-y``, ``-sses``/``-xes``/``-ches``/``-shes`` -> drop ``es``,
    otherwise drop a trailing ``s``).
    """
    head, _, last = table_name.rpartition("_")
    if last.endswith("ies"):
        last = last[:-3] + "y"
    elif last.endswith(("sses", "xes", "ches", "shes")):
        last = last[:-2]
    elif last.endswith("s") and not last.endswith("ss"):
        last = last[:-1]
    return f"{head}_{last}" if head else last


def _lower_first(text: str) -> str:
    """Lower-case the first character (for "Lịch sử tổ chức" from "Tổ chức")."""
    return text[:1].lower() + text[1:]


def _build_history_table(
    source: TableInfo, tables: dict[str, TableInfo]
) -> tuple[TableInfo, list[RefInfo]]:
    """Generate the history table of a tracked table, with its foreign keys.

    The change-history rule (``.claude/rules/database.md``): a history row is
    a copy of the *whole* old row, so the history table holds every source
    column with the same type - nullable, without primary-key/unique flags,
    because one source row appears many times and a column added later has
    no value in older rows - plus ``HISTORY_METADATA_COLUMNS``. History tables
    are never written in the DBML; this is their only definition.

    Returns:
        The history ``TableInfo`` and its refs (source key, ``changed_by``).
    """
    name = f"{_singular(source.name)}_history"
    users = tables.get("users")
    user_example = (
        next((column.example for column in users.columns if column.is_pk), "")
        if users
        else ""
    )
    key_columns = [column.name for column in source.columns if column.is_pk]
    columns = [
        ColumnInfo(
            name="history_id",
            type_label="bigint",
            enum_values=None,
            is_pk=True,
            is_not_null=True,
            is_unique=False,
            meaning="Auto-increasing ID of the history row.",
            vi_meaning="Mã tự tăng của dòng lịch sử.",
            example="1024",
            is_planned=False,
            planned_detail="",
        )
    ]
    for column in source.columns:
        # The history table is designed for the target: what the refactor
        # drops is never copied.
        if column.is_removed:
            continue
        enum_values = column.enum_values
        if enum_values is not None:
            enum_values = tuple(
                value
                for value in enum_values
                if value not in column.removed_enum_values
            )
        columns.append(
            ColumnInfo(
                name=column.name,
                type_label=column.type_label,
                enum_values=enum_values,
                is_pk=False,
                is_not_null=False,
                is_unique=False,
                meaning=f"Value before the change ({source.name}.{column.name}).",
                vi_meaning=(
                    f"Giá trị trước khi thay đổi ({source.name}.{column.name})."
                ),
                example=column.example,
                is_planned=False,
                planned_detail="",
            )
        )
    columns += [
        ColumnInfo(
            name="changed_at",
            type_label="timestamptz",
            enum_values=None,
            is_pk=False,
            is_not_null=True,
            is_unique=False,
            meaning="When this version of the row was replaced.",
            vi_meaning="Thời điểm phiên bản này của dòng bị thay thế.",
            example="2026-09-10T07:15:00Z",
            is_planned=False,
            planned_detail="",
        ),
        ColumnInfo(
            name="changed_by",
            type_label="uuid",
            enum_values=None,
            is_pk=False,
            is_not_null=False,
            is_unique=False,
            meaning="User who made the change; NULL when the system made it.",
            vi_meaning=(
                "Người dùng đã thực hiện thay đổi; NULL nếu do hệ thống thực hiện."
            ),
            example=user_example,
            is_planned=False,
            planned_detail="",
        ),
        ColumnInfo(
            name="change_reason",
            type_label="varchar(200)",
            enum_values=None,
            is_pk=False,
            is_not_null=True,
            is_unique=False,
            meaning=(
                "Why the row was changed, set by the application for the "
                "transaction: typed by the person for an administrative "
                "decision, a fixed text for a routine action. A change "
                "without a reason fails."
            ),
            vi_meaning=(
                "Lý do thay đổi dòng, do ứng dụng đặt cho giao dịch: người "
                "dùng nhập với quyết định quản trị, văn bản cố định với thao "
                "tác thường lệ. Thay đổi không có lý do sẽ bị từ chối."
            ),
            example="Customer moved to a new office",
            is_planned=False,
            planned_detail="",
        ),
    ]
    history = TableInfo(
        name=name,
        domain=source.domain,
        # A history table of a built table is itself not built until the
        # bulk refactor; otherwise it shares its source's design status.
        status="planned" if source.status == "built" else source.status,
        owner=source.owner,
        features=list(source.features),
        hypertable_column="",
        description=(
            f"Every earlier version of a row of `{source.name}`: a copy of the "
            "whole row, taken just before a change and written by a database "
            "trigger in the same transaction. Generated by the domain-model "
            "tool from `@tracked *`; never written by hand."
        ),
        vi_name=f"Lịch sử {_lower_first(source.vi_name)}" if source.vi_name else "",
        vi_description=(
            f"Mọi phiên bản trước đây của một dòng trong bảng {source.name}: bản "
            "sao toàn bộ dòng, chụp ngay trước khi thay đổi và được trigger ghi "
            "trong cùng giao dịch. Do công cụ domain-model tự sinh từ "
            "@tracked *, không bao giờ viết tay."
        ),
        columns=columns,
        indexes=[
            IndexInfo(
                label=", ".join([*key_columns, "changed_at"]),
                name=f"ix_{name}_{'_'.join(key_columns)}_time",
                is_unique=False,
                note="",
            )
        ],
        kind="history",
        parent=source.name,
    )
    refs: list[RefInfo] = []
    if len(key_columns) == 1:
        refs.append(
            RefInfo(
                from_table=name,
                from_column=key_columns[0],
                to_table=source.name,
                to_column=key_columns[0],
                kind="many-to-one",
                on_delete="restrict",
                is_planned=True,
            )
        )
    if users:
        user_key = next(column.name for column in users.columns if column.is_pk)
        refs.append(
            RefInfo(
                from_table=name,
                from_column="changed_by",
                to_table="users",
                to_column=user_key,
                kind="many-to-one",
                on_delete="restrict",
                is_planned=True,
            )
        )
    return history, refs


def _link_companions(
    tables: dict[str, TableInfo],
    domains: list[DomainInfo],
    refs: list[RefInfo],
    state_parents: dict[str, str],
) -> list[str]:
    """Generate history tables, link state tables, then number every table.

    Main and state tables share one sequence, 1, 2, ..., in design order
    (domain order, then table order inside the domain); a state table takes
    the number right after its main table. A generated history table is
    ``<N>.h`` of its source, listed right after it.

    Args:
        tables: Parsed DBML tables; generated history tables are added.
        domains: Domains; their ``table_names`` are reordered and extended.
        refs: Foreign keys; the history tables' refs are appended.
        state_parents: ``{state table: main table}`` from ``@state-of`` tags.

    Returns:
        Convention errors.
    """
    errors: list[str] = []
    for state_name, parent_name in state_parents.items():
        parent = tables.get(parent_name)
        state = tables[state_name]
        if parent is None:
            errors.append(f"{state_name}: @state-of names unknown table {parent_name}")
            continue
        if parent.state_table:
            errors.append(f"{parent_name} has two state tables")
        if parent.domain != state.domain:
            errors.append(f"{state_name} must be in the same domain as {parent_name}")
        parent.state_table, state.kind, state.parent = state_name, "state", parent_name
    for source in list(tables.values()):
        if not source.tracked_columns:
            continue
        for name in source.tracked_columns:
            if source.column(name) is None:
                errors.append(f"{source.name}: @tracked names unknown column {name}")
        history, history_refs = _build_history_table(source, tables)
        if history.name in tables:
            errors.append(f"{history.name} is generated; remove it from the DBML")
            continue
        tables[history.name] = history
        source.history_table = history.name
        refs += history_refs
    main_number = 0
    for domain in domains:
        ordered: list[str] = []
        for table_name in domain.table_names:
            table = tables.get(table_name)
            if table is None or table.kind != "main":
                continue
            main_number += 1
            table.number = str(main_number)
            ordered.append(table_name)
            # A state table is designed and reviewed like any table, so it
            # takes the next number; only the generated history is "<N>.h".
            if table.history_table:
                tables[table.history_table].number = f"{main_number}.h"
                ordered.append(table.history_table)
            if table.state_table:
                main_number += 1
                tables[table.state_table].number = str(main_number)
                ordered.append(table.state_table)
        domain.table_names = ordered
    return errors


def load_domain_model(source_path: Path) -> DomainModel:
    """Parse and validate the DBML source.

    Args:
        source_path: Path to ``domain-model.dbml``.

    Returns:
        The parsed model, with non-fatal findings in ``warnings``.

    Raises:
        DomainModelError: When the source breaks a convention (missing or
            invalid tags, a table outside every group or in two groups, a
            built column referencing an unbuilt table without ``@planned``).
    """
    database = PyDBML(source_path.read_text(encoding="utf-8"))
    errors: list[str] = []

    table_domains: dict[str, str] = {}
    domains: list[DomainInfo] = []
    for table_group in database.table_groups:
        table_names = [table.name for table in table_group.items]
        for table_name in table_names:
            if table_name in table_domains:
                errors.append(
                    f"table {table_name} is in two groups: "
                    f"{table_domains[table_name]}, {table_group.name}"
                )
            table_domains[table_name] = table_group.name
        group_note, vi_title, vi_note = _split_vietnamese(_note_text(table_group.note))
        domains.append(
            DomainInfo(
                name=table_group.name,
                note=group_note,
                vi_title=vi_title,
                vi_note=vi_note,
                table_names=table_names,
            )
        )

    tables: dict[str, TableInfo] = {}
    state_parents: dict[str, str] = {}
    for pydbml_table in database.tables:
        table_name = pydbml_table.name
        note = _note_text(pydbml_table.note)
        tag_line, _, description = note.partition("\n")
        tags = _parse_tags(tag_line) if tag_line.startswith("@") else {}
        if not tags:
            description = note
        status = tags.get("status", "")
        owner = tags.get("owner", "")
        if status not in STATUSES:
            errors.append(f"table {table_name}: @status must be one of {STATUSES}")
        if owner not in OWNERS:
            errors.append(f"table {table_name}: @owner must be one of {OWNERS}")
        if table_name not in table_domains:
            errors.append(f"table {table_name} is not in any TableGroup")
        features = [
            code for code in re.split(r"[,\s]+", tags.get("features", "")) if code
        ]
        description, vi_name, vi_description = _split_vietnamese(description)
        columns = [_parse_column(column) for column in pydbml_table.columns]
        if status != "built":
            errors += [
                f"{table_name}.{column.name}: {REMOVED_TAG} is only for built "
                "tables; delete the column from the design instead"
                for column in columns
                if column.is_removed or column.removed_enum_values
            ]
        tracked_columns, tracking_errors = _resolve_tracked_columns(
            table_name, tags, columns
        )
        errors += tracking_errors
        if "history-of" in tags:
            errors.append(
                f"{table_name}: history tables are generated from @tracked *; "
                "remove the hand-written table"
            )
        if "state-of" in tags:
            state_parents[table_name] = tags["state-of"]
        tables[table_name] = TableInfo(
            name=table_name,
            domain=table_domains.get(table_name, ""),
            status=status,
            owner=owner,
            features=features,
            hypertable_column=tags.get("hypertable", ""),
            description=description,
            vi_name=vi_name,
            vi_description=vi_description,
            columns=columns,
            indexes=[_parse_index(index) for index in pydbml_table.indexes],
            tracked_columns=tracked_columns,
        )

    refs: list[RefInfo] = []
    for pydbml_ref in database.refs:
        if len(pydbml_ref.col1) != 1 or len(pydbml_ref.col2) != 1:
            errors.append("composite foreign keys are not supported by the generator")
            continue
        left_column, right_column = pydbml_ref.col1[0], pydbml_ref.col2[0]
        # Normalize so "from" is the FK holder: '>' and '-' are written FK
        # side first; '<' is the mirror of '>'.
        if pydbml_ref.type == "<":
            left_column, right_column = right_column, left_column
        kind = {">": "many-to-one", "<": "many-to-one", "-": "one-to-one"}.get(
            pydbml_ref.type, "many-to-many"
        )
        from_table = tables[left_column.table.name]
        to_table = tables[right_column.table.name]
        from_column = from_table.column(left_column.name)
        is_planned = (
            from_table.status != "built"
            or to_table.status != "built"
            or (from_column is not None and from_column.is_planned)
        )
        if (
            from_table.status == "built"
            and to_table.status != "built"
            and from_column is not None
            and not from_column.is_planned
        ):
            errors.append(
                f"{from_table.name}.{left_column.name} is a built column referencing "
                f"unbuilt table {to_table.name}; mark it {PLANNED_COLUMN_TAG}"
            )
        refs.append(
            RefInfo(
                from_table=from_table.name,
                from_column=left_column.name,
                to_table=to_table.name,
                to_column=right_column.name,
                kind=kind,
                on_delete=(pydbml_ref.on_delete or "").lower(),
                is_planned=is_planned,
            )
        )

    errors += _link_companions(tables, domains, refs, state_parents)

    if errors:
        raise DomainModelError("\n".join(errors))

    # DM-24: a customer-owned row carries organization_id only as its own
    # owner or as the owner at the time it was recorded; otherwise it reads
    # the organization through a parent, so it needs a link to a
    # customer-owned table.
    customer_parents = {
        ref.from_table
        for ref in refs
        if tables[ref.to_table].owner == "customer" and ref.to_table != ref.from_table
    }
    warnings = [
        f"customer-owned table {table.name} has neither {TENANT_COLUMN} nor a "
        "link to a customer-owned parent"
        for table in tables.values()
        if table.owner == "customer"
        and table.column(TENANT_COLUMN) is None
        and table.name not in customer_parents
        and table.kind != "history"
    ]
    documentation_gaps = [
        f"{table.name}.{column.name}: missing {what}"
        for table in tables.values()
        for column in table.columns
        for what, value in (
            ("meaning", column.meaning),
            ("Vietnamese meaning (@vi)", column.vi_meaning),
            ("example", column.example),
        )
        if not value
    ]
    documentation_gaps += [
        f"table {table.name}: missing {what}"
        for table in tables.values()
        for what, value in (
            ("Vietnamese name (@vi-name)", table.vi_name),
            ("Vietnamese description (@vi)", table.vi_description),
        )
        if not value
    ]
    documentation_gaps += [
        f"domain {domain.name}: missing {what}"
        for domain in domains
        for what, value in (
            ("Vietnamese title (@vi-name)", domain.vi_title),
            ("Vietnamese note (@vi)", domain.vi_note),
        )
        if not value
    ]
    sticky_notes = [
        (note.name, *_split_vietnamese(_note_text(note))[::2])
        for note in database.sticky_notes
    ]
    documentation_gaps += [
        f"note {note_name}: missing Vietnamese text (@vi)"
        for note_name, _, vi_text in sticky_notes
        if not vi_text
    ]
    return DomainModel(
        project_note=_note_text(database.project.note) if database.project else "",
        domains=domains,
        tables=tables,
        refs=refs,
        sticky_notes=sticky_notes,
        warnings=warnings,
        documentation_gaps=documentation_gaps,
    )


# ---------------------------------------------------------------------------
# Rendering helpers
# ---------------------------------------------------------------------------


def _md_cell(text: str) -> str:
    """Escape text for a single Markdown table cell."""
    return text.replace("|", "\\|").replace("\n", "<br>")


def _mermaid_type(type_label: str) -> str:
    """Reduce a SQL type to a token Mermaid accepts as an attribute type."""
    return re.sub(r"[^A-Za-z0-9_]+", "_", type_label).strip("_") or "unknown"


def _domain_status(domain: DomainInfo, tables: dict[str, TableInfo]) -> str:
    """Return ``built``, ``planned`` or ``partial`` for a whole domain."""
    statuses = {tables[name].status for name in domain.table_names}
    if statuses == {"built"}:
        return "built"
    if "built" not in statuses:
        return "planned"
    return "partial"


def _kind_counts(tables: dict[str, TableInfo]) -> dict[str, int]:
    """Count tables per kind: ``main``, ``history`` (generated), ``state``."""
    counts = {"main": 0, "history": 0, "state": 0}
    for table in tables.values():
        counts[table.kind] += 1
    return counts


def _status_counts(
    table_names: list[str], tables: dict[str, TableInfo]
) -> dict[str, int]:
    """Count tables per status."""
    counts = dict.fromkeys(STATUSES, 0)
    for name in table_names:
        counts[tables[name].status] += 1
    return counts


def _table_link(table_name: str, model: DomainModel, from_domain: str | None) -> str:
    """Markdown link to a table's L3 section, relative to the linking page.

    Args:
        table_name: Target table.
        model: The domain model.
        from_domain: Domain page the link is written on, or None for the
            overview page.
    """
    target_domain = model.tables[table_name].domain
    if from_domain is None:
        return f"[{table_name}](domains/{target_domain}.md#{table_name})"
    if target_domain == from_domain:
        return f"[{table_name}](#{table_name})"
    return f"[{table_name}]({target_domain}.md#{table_name})"


def _er_relationship(ref: RefInfo, model: DomainModel) -> str:
    """Render one foreign key as a Mermaid erDiagram relationship line."""
    fk_column = model.tables[ref.from_table].column(ref.from_column)
    is_optional = fk_column is None or not fk_column.is_not_null
    parent_marker = "o|" if is_optional else "||"
    child_marker = {"many-to-one": "}o", "one-to-one": "|o"}.get(ref.kind, "}o")
    if ref.kind == "many-to-many":
        parent_marker = "o{"
    line_style = ".." if ref.is_planned else "--"
    return (
        f"  {ref.from_table} {child_marker}{line_style}{parent_marker} "
        f'{ref.to_table} : "{ref.from_column}"'
    )


# ---------------------------------------------------------------------------
# L1: overview
# ---------------------------------------------------------------------------


def render_overview(model: DomainModel) -> str:
    """Render the L1 overview page (domain map, ownership, decisions)."""
    tables = model.tables
    all_counts = _status_counts(list(tables), tables)
    kind_counts = _kind_counts(tables)
    lines = [GENERATED_HEADER, "# Domain model: overview", "", model.project_note, ""]

    lines += [
        "## At a glance",
        "",
        f"**{len(tables)} tables in {len(model.domains)} domains:** "
        f"{kind_counts['main']} main, {kind_counts['history']} history (N.h, "
        f"generated), {kind_counts['state']} state · "
        f"{all_counts['built']} built, {all_counts['planned']} planned, "
        f"{all_counts['proposed']} proposed.",
        "",
    ]

    # Domain map: one node per domain, one arrow per pair of domains linked
    # by at least one foreign key (FK holder -> referenced domain). An arrow
    # is solid when at least one of those links is already built.
    lines += ["## Domain map", "", "```mermaid", "flowchart LR"]
    for domain in model.domains:
        counts = _status_counts(domain.table_names, tables)
        count_label = " · ".join(
            f"{counts[status]} {status}" for status in STATUSES if counts[status]
        )
        lines.append(
            f'  {domain.name}["{domain.title}<br/>{count_label}"]'
            f":::{_domain_status(domain, tables)}"
        )
    domain_edges: dict[tuple[str, str], bool] = {}
    tenant_linked_domains: list[str] = []
    for ref in model.refs:
        edge = (tables[ref.from_table].domain, tables[ref.to_table].domain)
        if ref.to_table == TENANT_TABLE:
            if edge[0] != edge[1] and edge[0] not in tenant_linked_domains:
                tenant_linked_domains.append(edge[0])
            continue
        if edge[0] != edge[1]:
            domain_edges[edge] = domain_edges.get(edge, False) or not ref.is_planned
    for (from_domain, to_domain), is_built in domain_edges.items():
        arrow = "-->" if is_built else "-.->"
        lines.append(f"  {from_domain} {arrow} {to_domain}")
    lines += [DOMAIN_MAP_CLASSES.rstrip("\n"), "```", ""]
    lines += [
        "Green = fully built · orange = partly built · grey dashed = not built yet. "
        "An arrow **A → B** means A's data points to B's; solid = at least one "
        "such link is built, dashed = all planned.",
        "",
    ]
    if tenant_linked_domains and TENANT_TABLE in tables:
        tenant_domain = model.tables[TENANT_TABLE].domain
        lines += [
            f"Not drawn, to keep the map readable: {len(tenant_linked_domains)} domains "
            f"also point to **{tenant_domain}** through `{TENANT_COLUMN}` "
            f"({TENANT_TABLE}): {', '.join(tenant_linked_domains)}.",
            "",
        ]

    lines += [
        "## Domains",
        "",
        "| Domain | Status | ✅ | 📋 | 🆕 | Tables |",
        "|---|---|---|---|---|---|",
    ]
    for domain in model.domains:
        counts = _status_counts(domain.table_names, tables)
        table_list = ", ".join(
            f"{tables[name].number} {name}" for name in domain.table_names
        )
        lines.append(
            f"| [{domain.title}](domains/{domain.name}.md) "
            f"| {_domain_status(domain, tables)} | {counts['built']} "
            f"| {counts['planned']} | {counts['proposed']} | {table_list} |"
        )
    lines.append("")

    lines += ["## Data ownership", "", "| Owner | Meaning | Tables |", "|---|---|---|"]
    for owner in OWNERS:
        owned = [
            _table_link(name, model, None)
            for domain in model.domains
            for name in domain.table_names
            if tables[name].owner == owner
        ]
        if owned:
            lines.append(
                f"| **{owner}** | {OWNER_MEANINGS[owner]} | {', '.join(owned)} |"
            )
    lines.append("")

    for note_name, note_text, _ in model.sticky_notes:
        lines += [f"## {note_name.replace('_', ' ').capitalize()}", "", note_text, ""]

    return "\n".join(lines).rstrip("\n") + "\n"


# ---------------------------------------------------------------------------
# L2 + L3: one page per domain
# ---------------------------------------------------------------------------


def render_domain_page(domain: DomainInfo, model: DomainModel) -> str:
    """Render one domain page: L2 entity diagram and L3 table details."""
    tables = model.tables
    domain_table_names = set(domain.table_names)
    counts = _status_counts(domain.table_names, tables)
    lines = [
        GENERATED_HEADER,
        f"# {domain.title}",
        "",
        "[← Overview](../overview.md)",
        "",
        " · ".join(
            f"{STATUS_BADGES[status]}: {counts[status]}"
            for status in STATUSES
            if counts[status]
        ),
        "",
        domain.note,
        "",
    ]

    # L2 diagram: this domain's tables with their key columns, plus any
    # table from another domain it links to (shown without columns).
    domain_refs = [
        ref
        for ref in model.refs
        if ref.from_table in domain_table_names or ref.to_table in domain_table_names
    ]
    external_tables = sorted(
        (
            {ref.from_table for ref in domain_refs}
            | {ref.to_table for ref in domain_refs}
        )
        - domain_table_names
    )
    fk_columns = {(ref.from_table, ref.from_column) for ref in model.refs}
    lines += ["## Diagram", "", "```mermaid", "erDiagram"]
    for table_name in domain.table_names:
        table = tables[table_name]
        lines.append(f"  {table_name} {{")
        for column in table.columns:
            keys = [
                key
                for key, flag in (
                    ("PK", column.is_pk),
                    ("FK", (table_name, column.name) in fk_columns),
                )
                if flag
            ]
            if not keys:
                continue
            comment = ' "planned"' if column.is_planned else ""
            lines.append(
                f"    {_mermaid_type(column.type_label)} {column.name} "
                f"{', '.join(keys)}{comment}"
            )
        lines.append("  }")
    lines += [_er_relationship(ref, model) for ref in domain_refs]
    lines += ["```", ""]
    lines.append(
        "Only key columns are shown. Solid line = built link, dashed = planned. "
        + (
            "Tables from other domains (no columns): "
            + ", ".join(
                _table_link(name, model, domain.name) for name in external_tables
            )
            + "."
            if external_tables
            else ""
        )
    )
    lines.append("")

    # L3: every table in full.
    lines += ["## Tables", ""]
    for table_name in domain.table_names:
        lines += _render_table_section(tables[table_name], domain, model)

    return "\n".join(lines).rstrip("\n") + "\n"


def _render_table_section(
    table: TableInfo, domain: DomainInfo, model: DomainModel
) -> list[str]:
    """Render the L3 section of one table (columns, enums, indexes, references)."""
    outgoing_refs = {
        ref.from_column: ref for ref in model.refs if ref.from_table == table.name
    }
    incoming_refs = [ref for ref in model.refs if ref.to_table == table.name]

    badges = [
        f"**No. {table.number}**",
        STATUS_BADGES[table.status],
        f"owner: **{table.owner}**",
    ]
    if table.features:
        badges.append("features: " + ", ".join(table.features))
    if table.hypertable_column:
        badges.append(f"hypertable on `{table.hypertable_column}`")
    if table.kind in COMPANION_LABELS:
        badges.append(
            f"{COMPANION_LABELS[table.kind]} of "
            + _table_link(table.parent, model, domain.name)
        )
    if table.state_table:
        badges.append(
            "live state in " + _table_link(table.state_table, model, domain.name)
        )
    lines = [f"### {table.name}", "", " · ".join(badges), "", table.description, ""]
    if table.history_table:
        lines += [
            f"{TRACKED_MARK} = tracked column: a change to it copies the whole old "
            f"row into {_table_link(table.history_table, model, domain.name)}.",
            "",
        ]

    lines += [
        "| Column | Type | Null | Key | References | Meaning | Example |",
        "|---|---|---|---|---|---|---|",
    ]
    for column in table.columns:
        keys = []
        if column.is_pk:
            keys.append("PK")
        if column.name in outgoing_refs:
            keys.append("FK")
        if column.is_unique:
            keys.append("UQ")
        if column.name in table.tracked_columns:
            keys.append(TRACKED_MARK)
        reference = ""
        if column.name in outgoing_refs:
            ref = outgoing_refs[column.name]
            reference = (
                f"{_table_link(ref.to_table, model, domain.name)}.{ref.to_column}"
            )
            if ref.on_delete:
                reference += f" (on delete {ref.on_delete})"
        meaning = column.meaning
        if column.is_planned:
            detail = f" ({column.planned_detail})" if column.planned_detail else ""
            meaning = f"**📋 planned{detail}**" + (f": {meaning}" if meaning else "")
        if column.is_removed:
            detail = f" ({column.removed_detail})" if column.removed_detail else ""
            meaning = f"**🗑️ to be removed{detail}**" + (
                f": {meaning}" if meaning else ""
            )
        example = f"`{_md_cell(column.example)}`" if column.example else ""
        lines.append(
            f"| `{column.name}` | {_md_cell(column.type_label)} "
            f"| {'no' if column.is_not_null else 'yes'} | {' '.join(keys)} "
            f"| {reference} | {_md_cell(meaning)} | {example} |"
        )
    lines.append("")

    enum_columns = [column for column in table.columns if column.enum_values]
    if enum_columns:
        lines += ["**Enum values**", ""]
        seen_enums: set[str] = set()
        for column in enum_columns:
            if column.type_label in seen_enums:
                continue
            seen_enums.add(column.type_label)
            values = [
                f"~~{value}~~ (to be removed)"
                if value in column.removed_enum_values
                else f"{value} (📋 planned)"
                if value in column.planned_enum_values
                else value
                for value in column.enum_values or ()
            ]
            lines.append(f"- `{column.type_label}`: {', '.join(values)}")
        lines.append("")

    if table.indexes:
        lines += ["**Indexes**", ""]
        for index in table.indexes:
            parts = [
                f"`{index.name}`" if index.name else "(unnamed)",
                f"({index.label})",
            ]
            if index.is_unique:
                parts.append("unique")
            if index.note:
                parts.append(f"- {index.note}")
            lines.append("- " + " ".join(parts))
        lines.append("")

    if incoming_refs:
        lines += ["**Referenced by**", ""]
        for ref in incoming_refs:
            planned = " (planned)" if ref.is_planned else ""
            lines.append(
                f"- {_table_link(ref.from_table, model, domain.name)}"
                f".{ref.from_column}{planned}"
            )
        lines.append("")
    return lines


def render_all(model: DomainModel) -> dict[Path, str]:
    """Render every generated file; returns ``{path: content}``."""
    rendered = {OVERVIEW_PATH: render_overview(model)}
    for domain in model.domains:
        rendered[DOMAINS_DIR / f"{domain.name}.md"] = render_domain_page(domain, model)
    return rendered


def _stale_generated_files(rendered: dict[Path, str]) -> list[Path]:
    """Return generated domain pages whose domain no longer exists."""
    if not DOMAINS_DIR.exists():
        return []
    return [
        path
        for path in DOMAINS_DIR.glob("*.md")
        if path not in rendered
        and path.read_text(encoding="utf-8").startswith(GENERATED_MARKER)
    ]


# ---------------------------------------------------------------------------
# Excel workbook: Overview sheet + one sheet per table
# ---------------------------------------------------------------------------

XLSX_PATH = MODEL_DIR / "domain-model.xlsx"
# The workbook is for business readers (BOD, BA), so every label and every
# text in it is Vietnamese: labels come from the constants below, content
# from the DBML's @vi / @vi-name tags. It deliberately shows no build
# progress (built/planned/proposed): progress is tracked elsewhere.
OVERVIEW_SHEET = "Tổng quan"
# Excel refuses sheet names longer than 31 characters; long table names are
# shortened with these word replacements first, then truncated. The full
# table name is always in cell A1 of its sheet.
SHEET_NAME_LIMIT = 31
SHEET_NAME_ABBREVIATIONS = (("configuration", "config"), ("measurements", "meas"))
XLSX_FONT = "Arial"
XLSX_HEADER_FILL = "1F4E78"
XLSX_DOMAIN_FILL = "D9E1F2"
XLSX_LINK_COLOR = "0563C1"
XLSX_MUTED_COLOR = "595959"
XLSX_OPEN_LINK_LABEL = "→ Mở"
VI_OWNER_LABELS = {
    "customer": (
        "Khách hàng",
        "Thuộc về một tổ chức (thường là khách hàng; công ty G3 cũng có thể sở hữu); có cột organization_id hoặc lấy qua bảng cha (DM-24).",
    ),
    "internal": (
        "Nội bộ",
        "Dữ liệu của chính chúng ta (đơn vị vận hành nền tảng), dùng chung cho mọi tổ chức.",
    ),
    "two-party": (
        "Hai bên",
        "Tài sản của G3 được khách hàng sử dụng; cả hai bên cùng liên quan.",
    ),
    "undecided": (
        "Chưa quyết định",
        "Đang chờ một quyết định còn mở (xem cuối trang).",
    ),
}
VI_ON_DELETE_LABELS = {
    "restrict": "chặn xóa",
    "cascade": "xóa theo",
    "set null": "đặt về NULL",
    "no action": "không làm gì",
}
OVERVIEW_HEADERS = (
    ("STT", 6),
    ("Miền nghiệp vụ", 30),
    ("Bảng", 34),
    ("Tên tiếng Việt", 28),
    ("Chủ sở hữu dữ liệu", 16),
    ("Tính năng", 16),
    ("Mô tả", 80),
    ("Số cột", 8),
    ("Mở sheet", 10),
)
TABLE_SHEET_HEADERS = (
    ("#", 5),
    ("Cột", 26),
    ("Kiểu dữ liệu", 22),
    ("Cho phép rỗng", 12),
    ("Khóa", 8),
    ("Tham chiếu tới", 40),
    ("Ý nghĩa", 60),
    ("Ví dụ", 42),
    ("Lưu lịch sử", 11),
)
# Legend rows on the Overview sheet: (term, meaning).
VI_LEGEND = (
    (
        "Miền nghiệp vụ (domain)",
        "Một phần của hệ thống sở hữu dữ liệu riêng, ví dụ Xe, Trạm sạc, Thanh toán. Mỗi miền gồm một hoặc nhiều bảng.",
    ),
    (
        "Bảng",
        "Một loại dữ liệu được lưu, ví dụ Xe hoặc Phiên sạc. Mỗi bảng có một sheet riêng liệt kê từng cột.",
    ),
    (
        "Tính năng",
        "Mã tính năng trong danh sách tính năng (feature-list.md), ví dụ F-B2.",
    ),
    (
        "PK / FK / UQ",
        "PK: khóa chính (định danh duy nhất của dòng). FK: khóa ngoại (trỏ tới dòng của bảng khác). UQ: giá trị không được trùng.",
    ),
    (
        "Ví dụ",
        "Giá trị mẫu. Các mã ví dụ khớp nhau giữa các bảng, nên có thể lần theo một bản ghi mẫu qua nhiều sheet.",
    ),
    ("NULL", "Ô không có giá trị. Ý nghĩa của NULL được giải thích trong cột Ý nghĩa."),
    (
        f"{TRACKED_MARK} Lưu lịch sử",
        "Cột được theo dõi: khi cột này thay đổi, toàn bộ dòng cũ được chép sang "
        "bảng lịch sử (tên kết thúc bằng _history) kèm thời điểm và người thay đổi.",
    ),
)


def _plain(text: str) -> str:
    """Strip the Markdown used in notes (bold, code, links) for Excel cells."""
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def _split_domain_note(note: str) -> tuple[str, list[str]]:
    """Split a domain note into its description and its relationship bullets."""
    description_lines: list[str] = []
    bullets: list[str] = []
    for line in note.splitlines():
        stripped = line.strip()
        if stripped.startswith("- "):
            bullets.append(_plain(stripped[2:]))
        elif stripped and not bullets:
            description_lines.append(_plain(stripped))
    return " ".join(description_lines), bullets


def _markdown_table_rows(text: str) -> list[list[str]]:
    """Return the cells of every Markdown table row in ``text`` (header included)."""
    rows = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("|") and not set(stripped) <= set("|-: "):
            rows.append(
                [_plain(cell.strip()) for cell in stripped.strip("|").split("|")]
            )
    return rows


def _sheet_names(model: DomainModel) -> dict[str, str]:
    """Assign every table a unique Excel sheet name of at most 31 characters."""
    used_names = {OVERVIEW_SHEET.lower()}
    sheet_names: dict[str, str] = {}
    for domain in model.domains:
        for table_name in domain.table_names:
            sheet_name = table_name
            for long_word, short_word in SHEET_NAME_ABBREVIATIONS:
                if len(sheet_name) > SHEET_NAME_LIMIT:
                    sheet_name = sheet_name.replace(long_word, short_word)
            base_name = sheet_name = sheet_name[:SHEET_NAME_LIMIT]
            suffix_number = 2
            while sheet_name.lower() in used_names:
                suffix = f"~{suffix_number}"
                sheet_name = base_name[: SHEET_NAME_LIMIT - len(suffix)] + suffix
                suffix_number += 1
            used_names.add(sheet_name.lower())
            sheet_names[table_name] = sheet_name
    return sheet_names


class _SheetWriter:
    """Row-by-row writer for one worksheet with this workbook's styling.

    Attributes:
        worksheet: The openpyxl worksheet being written.
        row: Next row to write (1-based).
        widths: Column widths in characters, used to estimate wrapped row heights.
    """

    def __init__(self, worksheet: object, widths: list[int]) -> None:
        """Set column widths and start at row 1."""
        from openpyxl.utils import get_column_letter  # noqa: PLC0415

        self.worksheet = worksheet
        self.row = 1
        self.widths = widths
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = width

    def write(
        self,
        values: list[object],
        *,
        bold: bool = False,
        size: int = 10,
        color: str | None = None,
        fill: str | None = None,
        fills: dict[int, str] | None = None,
        links: dict[int, str] | None = None,
        wrap: bool = True,
    ) -> int:
        """Write one row and return its number.

        Args:
            values: Cell values from column A; None leaves a cell empty.
            bold: Bold font for the whole row.
            size: Font size in points.
            color: Font color (hex RGB) for the whole row.
            fill: Background fill (hex RGB) for every written cell.
            fills: Per-column fills ``{1-based column: hex RGB}``, overriding ``fill``.
            links: Internal hyperlinks ``{1-based column: sheet name}``.
            wrap: Wrap text and grow the row height to fit.
        """
        from openpyxl.styles import Alignment, Font, PatternFill  # noqa: PLC0415
        from openpyxl.worksheet.hyperlink import Hyperlink  # noqa: PLC0415

        line_count = 1
        for column_index, value in enumerate(values, start=1):
            # An empty string is saved as an empty cell and reads back as
            # None; write None up front so a round trip compares equal.
            if value == "":
                value = None
            cell = self.worksheet.cell(row=self.row, column=column_index, value=value)
            link_target = (links or {}).get(column_index)
            cell.font = Font(
                name=XLSX_FONT,
                size=size,
                bold=bold,
                color=XLSX_LINK_COLOR if link_target else color,
                underline="single" if link_target else None,
            )
            if link_target:
                # An internal "location" link, not an external target, so
                # Excel and LibreOffice both jump to the sheet in place.
                cell.hyperlink = Hyperlink(
                    ref=cell.coordinate, location=f"'{link_target}'!A1"
                )
            cell_fill = (fills or {}).get(column_index, fill)
            if cell_fill:
                cell.fill = PatternFill("solid", fgColor=cell_fill)
            cell.alignment = Alignment(wrap_text=wrap, vertical="top")
            if wrap and value is not None and column_index <= len(self.widths):
                # openpyxl cannot auto-fit, and Excel does not re-fit rows on
                # open, so estimate the wrapped line count from the width.
                chars_per_line = max(self.widths[column_index - 1] - 2, 1)
                cell_lines = sum(
                    max(1, -(-len(part) // chars_per_line))
                    for part in str(value).split("\n")
                )
                line_count = max(line_count, cell_lines)
        if line_count > 1:
            self.worksheet.row_dimensions[self.row].height = 13.5 * line_count
        self.row += 1
        return self.row - 1

    def header(self, labels: list[str]) -> int:
        """Write a white-on-blue header row and return its number."""
        return self.write(labels, bold=True, color="FFFFFF", fill=XLSX_HEADER_FILL)

    def blank(self) -> None:
        """Skip one row."""
        self.row += 1


def _vi_domain_title(domain: DomainInfo) -> str:
    """Vietnamese domain title with the technical name, e.g. ``Đội xe (fleet)``."""
    return f"{domain.vi_title or domain.title} ({domain.name})"


def _write_overview_sheet(
    worksheet: object, model: DomainModel, sheet_names: dict[str, str]
) -> None:
    """Fill the Overview sheet: domain/table index first, then legend, decisions.

    Only the title rows and the index header are frozen: freezing more makes
    the frozen area fill the window on a laptop screen, so nothing scrolls.
    """
    writer = _SheetWriter(worksheet, [width for _, width in OVERVIEW_HEADERS])
    kind_counts = _kind_counts(model.tables)
    writer.write(["G3 Network: Mô hình dữ liệu"], bold=True, size=16, wrap=False)
    writer.write(
        [
            "Tệp được tạo tự động từ domain-model.dbml. Không sửa trực tiếp tệp "
            "này; hãy sửa tệp .dbml rồi chạy lại lệnh generate."
        ],
        color=XLSX_MUTED_COLOR,
        wrap=False,
    )
    writer.write(
        [
            f"{len(model.tables)} bảng trong {len(model.domains)} miền nghiệp vụ: "
            f"{kind_counts['main']} bảng chính, {kind_counts['history']} bảng lịch "
            f"sử (N.h), {kind_counts['state']} bảng trạng thái. "
            f"Bấm vào tên bảng hoặc «{XLSX_OPEN_LINK_LABEL}» để mở sheet chi tiết "
            "của bảng; chú giải và các quyết định còn mở nằm ở cuối trang."
        ],
        bold=True,
        wrap=False,
    )
    writer.blank()

    header_row = writer.header([label for label, _ in OVERVIEW_HEADERS])
    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
    for domain in model.domains:
        description, relationships = _split_domain_note(domain.vi_note)
        domain_text = description
        if relationships:
            domain_text += "\n" + "\n".join(f"• {line}" for line in relationships)
        writer.write(
            [
                None,
                _vi_domain_title(domain),
                f"{len(domain.table_names)} bảng",
                None,
                None,
                None,
                domain_text,
            ],
            bold=True,
            fill=XLSX_DOMAIN_FILL,
        )
        for table_name in domain.table_names:
            table = model.tables[table_name]
            sheet_name = sheet_names[table_name]
            writer.write(
                [
                    table.number,
                    _vi_domain_title(domain),
                    table_name,
                    table.vi_name,
                    VI_OWNER_LABELS[table.owner][0],
                    ", ".join(table.features),
                    _plain(table.vi_description),
                    len(table.columns),
                    XLSX_OPEN_LINK_LABEL,
                ],
                links={3: sheet_name, 4: sheet_name, 9: sheet_name},
            )

    writer.blank()
    writer.header(["Chú giải", "Ý nghĩa"])
    for owner in OWNERS:
        label, meaning = VI_OWNER_LABELS[owner]
        writer.write([f"Chủ sở hữu dữ liệu: {label}", meaning])
    for term, meaning in VI_LEGEND:
        writer.write([term, meaning])

    for note_name, _, vi_text in model.sticky_notes:
        writer.blank()
        title = "Các quyết định còn mở" if note_name == "open_decisions" else note_name
        writer.write([title], bold=True, size=12, wrap=False)
        note_rows = _markdown_table_rows(vi_text)
        if note_rows:
            writer.header(note_rows[0])
            for note_row in note_rows[1:]:
                writer.write(list(note_row))
        else:
            writer.write([_plain(vi_text)])


def _write_table_sheet(
    worksheet: object,
    table: TableInfo,
    model: DomainModel,
    sheet_names: dict[str, str],
) -> None:
    """Fill one table's sheet: header, columns, enums, indexes, references."""
    writer = _SheetWriter(worksheet, [width for _, width in TABLE_SHEET_HEADERS])
    outgoing_refs = {
        ref.from_column: ref for ref in model.refs if ref.from_table == table.name
    }
    incoming_refs = [ref for ref in model.refs if ref.to_table == table.name]
    domain = next(domain for domain in model.domains if domain.name == table.domain)

    writer.write(
        [f"{table.number}. {table.name}: {table.vi_name}"],
        bold=True,
        size=14,
        wrap=False,
    )
    writer.write([f"← {OVERVIEW_SHEET}"], links={1: OVERVIEW_SHEET}, wrap=False)
    facts = [
        f"Miền nghiệp vụ: {_vi_domain_title(domain)}",
        f"Chủ sở hữu dữ liệu: {VI_OWNER_LABELS[table.owner][0]}",
    ]
    if table.features:
        facts.append("Tính năng: " + ", ".join(table.features))
    if table.hypertable_column:
        facts.append(f"Dữ liệu chuỗi thời gian theo cột {table.hypertable_column}")
    writer.write(["   ·   ".join(facts)], wrap=False)
    writer.write([_plain(table.vi_description).replace("\n", " ")], wrap=False)
    if table.history_table:
        writer.write(
            [
                f"{TRACKED_MARK} Lưu lịch sử: khi một cột có dấu {TRACKED_MARK} thay "
                f"đổi, toàn bộ dòng cũ được chép sang bảng {table.history_table}."
            ],
            links={1: sheet_names[table.history_table]},
            wrap=False,
        )
    if table.state_table:
        writer.write(
            [f"Trạng thái trực tiếp lưu ở bảng: {table.state_table}"],
            links={1: sheet_names[table.state_table]},
            wrap=False,
        )
    if table.kind in VI_COMPANION_LABELS:
        writer.write(
            [f"{VI_COMPANION_LABELS[table.kind]}: {table.parent}"],
            links={1: sheet_names[table.parent]},
            wrap=False,
        )
    writer.blank()

    header_row = writer.header([label for label, _ in TABLE_SHEET_HEADERS])
    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
    for position, column in enumerate(table.columns, start=1):
        keys = [
            key
            for key, is_set in (
                ("PK", column.is_pk),
                ("FK", column.name in outgoing_refs),
                ("UQ", column.is_unique),
            )
            if is_set
        ]
        reference = ""
        links: dict[int, str] = {}
        if column.name in outgoing_refs:
            ref = outgoing_refs[column.name]
            reference = f"{ref.to_table}.{ref.to_column}"
            if ref.on_delete:
                on_delete = VI_ON_DELETE_LABELS.get(ref.on_delete, ref.on_delete)
                reference += f" (khi xóa: {on_delete})"
            links[6] = sheet_names[ref.to_table]
        meaning = _plain(column.vi_meaning)
        # A planned column's decision IDs are not progress: they point the
        # reader to the open question the column depends on.
        decision_ids = re.findall(r"D\d+", column.planned_detail)
        if decision_ids and not all(code in meaning for code in decision_ids):
            meaning += f" (xem quyết định {', '.join(decision_ids)})"
        if column.is_removed:
            detail = f": {column.removed_detail}" if column.removed_detail else ""
            meaning = f"[Sẽ bỏ{detail}] {meaning}"
        writer.write(
            [
                position,
                column.name,
                column.type_label,
                "không" if column.is_not_null else "có",
                " ".join(keys),
                reference,
                meaning,
                column.example,
                TRACKED_MARK if column.name in table.tracked_columns else None,
            ],
            links=links,
        )
    worksheet.auto_filter.ref = f"A{header_row}:I{writer.row - 1}"

    enum_columns = [column for column in table.columns if column.enum_values]
    if enum_columns:
        writer.blank()
        writer.header(["", "Kiểu enum", "Giá trị cho phép"])
        seen_enums: set[str] = set()
        for column in enum_columns:
            if column.type_label not in seen_enums:
                seen_enums.add(column.type_label)
                values = [
                    f"{value} (sẽ bỏ)"
                    if value in column.removed_enum_values
                    else f"{value} (dự kiến)"
                    if value in column.planned_enum_values
                    else value
                    for value in column.enum_values or ()
                ]
                writer.write([None, column.type_label, ", ".join(values)])

    if table.indexes:
        writer.blank()
        writer.header(["", "Chỉ mục (index)", "Các cột", "Duy nhất", "", "Ghi chú"])
        for index in table.indexes:
            writer.write(
                [
                    None,
                    index.name or "(không tên)",
                    index.label,
                    "có" if index.is_unique else "không",
                    None,
                    index.note,
                ]
            )

    if incoming_refs:
        writer.blank()
        writer.header(["", "Được tham chiếu bởi bảng", "Cột"])
        for ref in incoming_refs:
            writer.write(
                [None, ref.from_table, ref.from_column],
                links={2: sheet_names[ref.from_table]},
            )


def build_workbook(model: DomainModel) -> object:
    """Build the Excel export: an Overview sheet, then one sheet per table.

    Tables appear in domain order, the same order as the Overview index.
    """
    from openpyxl import Workbook  # noqa: PLC0415

    workbook = Workbook()
    sheet_names = _sheet_names(model)
    overview_sheet = workbook.active
    overview_sheet.title = OVERVIEW_SHEET
    _write_overview_sheet(overview_sheet, model, sheet_names)
    for domain in model.domains:
        for table_name in domain.table_names:
            table_sheet = workbook.create_sheet(sheet_names[table_name])
            _write_table_sheet(
                table_sheet, model.tables[table_name], model, sheet_names
            )
    return workbook


def _workbook_values(workbook: object) -> list[tuple[str, list[tuple[object, ...]]]]:
    """Return every sheet's title and cell values, for comparing two workbooks.

    Byte comparison is useless for .xlsx (the zip holds timestamps), so
    ``generate`` and ``check`` compare values instead.
    """
    return [
        (sheet.title, [tuple(row) for row in sheet.iter_rows(values_only=True)])
        for sheet in workbook.worksheets
    ]


def _is_workbook_current(workbook: object) -> bool:
    """Whether the committed .xlsx holds the same values as ``workbook``."""
    from openpyxl import load_workbook  # noqa: PLC0415

    if not XLSX_PATH.exists():
        return False
    return _workbook_values(load_workbook(XLSX_PATH)) == _workbook_values(workbook)


# ---------------------------------------------------------------------------
# Check against the backend models
# ---------------------------------------------------------------------------

# SQL type spellings that mean the same thing in PostgreSQL.
TYPE_ALIASES = {
    "timestampwithtimezone": "timestamptz",
    "doubleprecision": "float8",
    "boolean": "bool",
    "integer": "int4",
    "int": "int4",
    "bigint": "int8",
}


def _normalize_type(type_label: str) -> str:
    """Normalize a SQL type spelling for comparison."""
    compact = type_label.lower().replace(" ", "").replace('"', "")
    return TYPE_ALIASES.get(compact, compact)


def _load_backend_metadata() -> object:
    """Import every domain's models module and return ``Base.metadata``.

    Side effects:
        Adds ``backend/`` to ``sys.path`` and imports all
        ``app.domains.*.models`` modules.
    """
    sys.path.insert(0, str(REPO_ROOT / "backend"))
    import app.domains as backend_domains  # noqa: PLC0415

    for module_info in pkgutil.iter_modules(backend_domains.__path__):
        try:
            importlib.import_module(f"app.domains.{module_info.name}.models")
        except ModuleNotFoundError as error:
            if error.name != f"app.domains.{module_info.name}.models":
                raise
    from app.libs.db.base import Base  # noqa: PLC0415

    return Base.metadata


def check_against_backend(model: DomainModel) -> list[str]:
    """Compare every built table with the SQLAlchemy metadata.

    Args:
        model: The parsed domain model.

    Returns:
        Human-readable mismatches; empty when the model matches the code.
    """
    from sqlalchemy import Enum as SqlEnum  # noqa: PLC0415
    from sqlalchemy.dialects import postgresql  # noqa: PLC0415

    dialect = postgresql.dialect()
    metadata = _load_backend_metadata()
    problems: list[str] = []
    refs_by_column = {
        (ref.from_table, ref.from_column): ref
        for ref in model.refs
        if not ref.is_planned
    }

    for sql_table in metadata.sorted_tables:
        table = model.tables.get(sql_table.name)
        if table is None:
            problems.append(
                f"{sql_table.name}: exists in the models but not in the DBML"
            )
        elif table.status != "built":
            problems.append(
                f"{sql_table.name}: exists in the models but is @status {table.status}"
            )

    for table in model.tables.values():
        if table.status != "built":
            continue
        sql_table = metadata.tables.get(table.name)
        if sql_table is None:
            problems.append(f"{table.name}: @status built but not found in the models")
            continue
        dbml_columns = {
            column.name: column for column in table.columns if not column.is_planned
        }
        for missing in sorted(set(sql_table.columns.keys()) - set(dbml_columns)):
            problems.append(
                f"{table.name}.{missing}: in the models, missing from the DBML"
            )
        for extra in sorted(set(dbml_columns) - set(sql_table.columns.keys())):
            problems.append(
                f"{table.name}.{extra}: in the DBML, not in the models "
                f"(mark it {PLANNED_COLUMN_TAG} if it is a proposal)"
            )
        for sql_column in sql_table.columns:
            column = dbml_columns.get(sql_column.name)
            if column is None:
                continue
            where = f"{table.name}.{sql_column.name}"
            if isinstance(sql_column.type, SqlEnum):
                if column.type_label != sql_column.type.name:
                    problems.append(
                        f"{where}: enum {sql_column.type.name} in the models, "
                        f"{column.type_label} in the DBML"
                    )
                elif tuple(sql_column.type.enums) != (
                    built_values := tuple(
                        value
                        for value in column.enum_values or ()
                        if value not in column.planned_enum_values
                    )
                ):
                    problems.append(
                        f"{where}: enum values differ: models {sql_column.type.enums}, "
                        f"DBML {list(built_values)} (without @planned values)"
                    )
            else:
                sql_type = _normalize_type(sql_column.type.compile(dialect=dialect))
                if sql_type != _normalize_type(column.type_label):
                    problems.append(
                        f"{where}: type {sql_type} in the models, "
                        f"{_normalize_type(column.type_label)} in the DBML"
                    )
            if sql_column.primary_key != column.is_pk:
                problems.append(f"{where}: primary key differs")
            if (not sql_column.nullable) != column.is_not_null:
                problems.append(
                    f"{where}: nullability differs (models nullable={sql_column.nullable})"
                )
            sql_foreign_keys = list(sql_column.foreign_keys)
            ref = refs_by_column.get((table.name, sql_column.name))
            if sql_foreign_keys and ref is None:
                problems.append(
                    f"{where}: foreign key to {sql_foreign_keys[0].target_fullname} "
                    "in the models, no built Ref in the DBML"
                )
            elif ref is not None and not sql_foreign_keys:
                problems.append(
                    f"{where}: Ref in the DBML, no foreign key in the models"
                )
            elif ref is not None:
                sql_foreign_key = sql_foreign_keys[0]
                if sql_foreign_key.target_fullname != f"{ref.to_table}.{ref.to_column}":
                    problems.append(
                        f"{where}: references {sql_foreign_key.target_fullname} in the "
                        f"models, {ref.to_table}.{ref.to_column} in the DBML"
                    )
                sql_on_delete = (sql_foreign_key.ondelete or "").lower()
                if sql_on_delete != ref.on_delete:
                    problems.append(
                        f"{where}: on delete {sql_on_delete or 'none'} in the models, "
                        f"{ref.on_delete or 'none'} in the DBML"
                    )
    return problems


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------


def command_generate(model: DomainModel) -> int:
    """Write every generated view and delete pages of removed domains."""
    rendered = render_all(model)
    DOMAINS_DIR.mkdir(parents=True, exist_ok=True)
    for path, content in rendered.items():
        path.write_text(content, encoding="utf-8")
        print(f"wrote {path.relative_to(REPO_ROOT)}")
    for path in _stale_generated_files(rendered):
        path.unlink()
        print(f"removed {path.relative_to(REPO_ROOT)}")
    # Rewrite the workbook only when its content changed: every save embeds
    # a fresh timestamp, which would otherwise make git see a new binary.
    workbook = build_workbook(model)
    if _is_workbook_current(workbook):
        print(f"unchanged {XLSX_PATH.relative_to(REPO_ROOT)}")
    else:
        workbook.save(XLSX_PATH)
        print(f"wrote {XLSX_PATH.relative_to(REPO_ROOT)}")
    return 0


def command_check(model: DomainModel) -> int:
    """Check the model against the backend and the generated views."""
    problems = check_against_backend(model) + model.documentation_gaps
    rendered = render_all(model)
    for path, content in rendered.items():
        if not path.exists() or path.read_text(encoding="utf-8") != content:
            problems.append(f"{path.relative_to(REPO_ROOT)}: stale, run generate")
    for path in _stale_generated_files(rendered):
        problems.append(f"{path.relative_to(REPO_ROOT)}: domain removed, run generate")
    if not _is_workbook_current(build_workbook(model)):
        problems.append(f"{XLSX_PATH.relative_to(REPO_ROOT)}: stale, run generate")
    for problem in problems:
        print(f"MISMATCH {problem}")
    built_count = sum(table.status == "built" for table in model.tables.values())
    if problems:
        print(f"{len(problems)} problem(s) found")
        return 1
    print(f"OK: {built_count} built tables match the models; views are up to date")
    return 0


def main() -> int:
    """Parse arguments, load the model, and run the chosen command."""
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("command", choices=("generate", "check"))
    arguments = parser.parse_args()
    try:
        model = load_domain_model(SOURCE_PATH)
    except DomainModelError as error:
        print(f"ERROR in {SOURCE_PATH.relative_to(REPO_ROOT)}:\n{error}")
        return 2
    for warning in model.warnings:
        print(f"WARNING {warning}")
    if arguments.command == "generate":
        for gap in model.documentation_gaps:
            print(f"WARNING {gap}")
        return command_generate(model)
    return command_check(model)


if __name__ == "__main__":
    sys.exit(main())
